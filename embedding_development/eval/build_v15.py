#!/usr/bin/env python3
"""
build_v15.py  —  Apply figure/text corrections to v14 (68 slides).

Fixes:
  1. S26 (idx 25): replace r2_comparison_grand_mean.png  (error bars removed)
  2. S35 (idx 34): replace trend_noise_comparison.png    (em-dashes → parentheses)
  3. S38 (idx 37): replace gpv_task_ensembles.png        (Cue/Choice label overlap fixed)
  4. S49 (idx 48): replace embedding_linear_map.png      (rotated labels, shorter arch names)
  5. S50 (idx 49): fix slide background to white         (was dark #1a1a2e)

Net: 68 slides (all in-place)
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR

base  = os.path.dirname(os.path.abspath(__file__))
root  = os.path.join(base, '..')
mdir  = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')
cdir  = os.path.join(root, 'outputs', 'cebra_comparison')
ml_dir = os.path.join(root, 'outputs', 'mlps', 'ml_vs_naive')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v14.pptx')
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v15.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation_v15.pptx'),
]

SLIDE_W = Inches(10.0)
SLIDE_H = Inches(5.625)


def replace_figure(slide, img_path):
    pics = [s for s in slide.shapes if s.shape_type == 13]
    if not pics:
        print(f'  WARNING: no picture on slide')
        return
    target = max(pics, key=lambda s: s.width * s.height)
    l, t, w, h = int(target.left), int(target.top), int(target.width), int(target.height)
    target._element.getparent().remove(target._element)
    slide.shapes.add_picture(img_path, l, t, w, h)


prs = Presentation(SRC)
print(f'Opened v14: {len(prs.slides)} slides')

# 1. S26 — grand mean R² bar (no error bars)
img = os.path.join(cdir, 'r2_comparison_grand_mean.png')
if os.path.exists(img):
    replace_figure(prs.slides[25], img)
    print('  S26: replaced r2_comparison_grand_mean.png')
else:
    print(f'  S26: SKIP ({img})')

# 2. S35 — trend noise (no em-dashes)
img = os.path.join(mdir, 'trend_noise_comparison.png')
if os.path.exists(img):
    replace_figure(prs.slides[34], img)
    print('  S35: replaced trend_noise_comparison.png')
else:
    print(f'  S35: SKIP ({img})')

# 3. S38 — GPV task ensembles (overlap fixed)
img = os.path.join(mdir, 'gpv_task_ensembles.png')
if os.path.exists(img):
    replace_figure(prs.slides[37], img)
    print('  S38: replaced gpv_task_ensembles.png')
else:
    print(f'  S38: SKIP ({img})')

# 4. S49 — embedding linear map (rotated labels)
img = os.path.join(mdir, 'embedding_linear_map.png')
if os.path.exists(img):
    replace_figure(prs.slides[48], img)
    print('  S49: replaced embedding_linear_map.png')
else:
    print(f'  S49: SKIP ({img})')

# 5. S50 — fix background to white
slide50 = prs.slides[49]
fill = slide50.background.fill
fill.solid()
fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
# Also update text colours from white → dark
for sh in slide50.shapes:
    try:
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if run.font.color.rgb == RGBColor(0xFF, 0xFF, 0xFF):
                    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                elif run.font.color.rgb == RGBColor(0xCC, 0xCC, 0xCC):
                    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    except Exception:
        pass
print('  S50: background set to white, text darkened')

# Save
print(f'\nFinal slide count: {len(prs.slides)}')
for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved  → {out}')
