#!/usr/bin/env python3
"""
build_v35.py  —  Fix cue labels on S65 (Cue A/B → Cue 1/2).

v34 (88) → v35 (88 slides, same count):
  S65: replace e07_joint_tuning.png (legend now shows Cue 1 / Cue 2)
"""
import os
from pptx import Presentation
from pptx.util import Inches
from pptx.parts.presentation import PresentationPart


@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
adir = os.path.join(root, 'outputs', 'ablation_vs_attribution')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v34.pptx')
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v35.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation_v35.pptx'),
]


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def replace_picture(slide, new_img, left_in, top_in, w_in, h_in):
    for sh in list(slide.shapes):
        if sh.shape_type == 13:
            sh._element.getparent().remove(sh._element)
            break
    slide.shapes.add_picture(new_img,
        int(Inches(left_in)), int(Inches(top_in)),
        int(Inches(w_in)),    int(Inches(h_in)))


prs = Presentation(SRC)
print(f'Opened v34: {len(prs.slides)} slides')

idx = find_idx(prs, 'joint position')
if idx is not None:
    replace_picture(prs.slides[idx],
                    os.path.join(adir, 'e07_joint_tuning.png'),
                    0.25, 0.68, 9.50, 4.20)
    print(f'  S{idx+1}: replaced e07_joint_tuning.png (Cue 1/2 labels)')

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved → {out}')
