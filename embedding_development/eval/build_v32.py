#!/usr/bin/env python3
"""
build_v32.py  —  Slide title fix + appendix AI-labelling + new A18 bucket-analysis slide.

v31 (86) → v32 (87 slides, +1):
  1. S62 title: remove "supervisor's variable" framing → "Neural signal: E07 × Cue, E23 × Choice"
  2. Regenerated figures already in place (bucket analysis, language cleaned)
  3. Add "AI-generated" badge (bottom-right) to all appendix slides (S69–S86)
  4. Update A02 (GPV) APPEARS ON to include S63; add GPV/R² note
  5. Update A09 (η²/Cohen's d) APPEARS ON; add bucket-analysis usage note
  6. Add new A18 slide: Bucket Analysis Methodology
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.parts.presentation import PresentationPart


@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v31.pptx')
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v32.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation_v32.pptx'),
]

C_BG    = RGBColor(0xFF, 0xFF, 0xFF)
C_TAG   = RGBColor(0x99, 0x99, 0x99)
C_TITLE = RGBColor(0x1A, 0x23, 0x3A)
C_HDR   = RGBColor(0x33, 0x33, 0x33)
C_BODY  = RGBColor(0x22, 0x22, 0x22)
C_LINE  = RGBColor(0xCC, 0xCC, 0xCC)


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def patch_run(slide, old_sub, new_sub):
    """Replace first occurrence of old_sub in any run on slide."""
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if old_sub in run.text:
                    run.text = run.text.replace(old_sub, new_sub)
                    return True
    return False


def patch_shape_text(slide, shape_idx, new_text):
    """Replace full text of the text frame in shape at shape_idx."""
    sh = slide.shapes[shape_idx]
    tf = sh.text_frame
    # Clear all paragraphs except the first, set first paragraph text
    for para in tf.paragraphs[1:]:
        for run in para.runs:
            run.text = ''
    if tf.paragraphs:
        runs = tf.paragraphs[0].runs
        if runs:
            runs[0].text = new_text
            for r in runs[1:]:
                r.text = ''
        else:
            from pptx.oxml.ns import qn
            from lxml import etree
            r_elm = etree.SubElement(tf.paragraphs[0]._p, qn('a:r'))
            rPr = etree.SubElement(r_elm, qn('a:rPr'), attrib={'lang': 'en-US', 'dirty': '0'})
            t = etree.SubElement(r_elm, qn('a:t'))
            t.text = new_text


def add_ai_badge(slide):
    """Add a small 'AI-generated' label at bottom-right of slide."""
    tb = slide.shapes.add_textbox(
        int(Inches(7.20)), int(Inches(5.27)),
        int(Inches(2.50)), int(Inches(0.22)))
    tf = tb.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    from pptx.enum.text import PP_ALIGN
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = 'AI-generated'
    r.font.size = Pt(7)
    r.font.italic = True
    r.font.color.rgb = C_TAG


def add_appendix_slide(prs, tag, title, col1_hdr, col1_body,
                       col2_hdr, col2_body, col3_hdr, col3_body):
    """Add one appendix slide matching the standard 3-column hidden layout."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill; fill.solid(); fill.fore_color.rgb = C_BG

    def tb(left, top, w, h, text, size, bold=False, color=C_BODY, italic=False):
        box = slide.shapes.add_textbox(
            int(Inches(left)), int(Inches(top)),
            int(Inches(w)), int(Inches(h)))
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; r = p.add_run()
        r.text = text; r.font.size = Pt(size)
        r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color

    tb(0.30, 0.05, 2.50, 0.28, '[APPENDIX — HIDDEN]', 8, color=C_TAG)
    tb(0.30, 0.28, 9.40, 0.52, title, 14, bold=True, color=C_TITLE)
    # thin rule
    from pptx.util import Emu
    from pptx.oxml.ns import qn
    from lxml import etree
    rule = slide.shapes.add_textbox(int(Inches(0.30)), int(Inches(0.83)),
                                    int(Inches(9.40)), int(Inches(0.01)))

    for col, (hdr, body) in enumerate([(col1_hdr, col1_body),
                                        (col2_hdr, col2_body),
                                        (col3_hdr, col3_body)]):
        x = 0.30 + col * 3.13
        tb(x, 0.92, 2.98, 0.28, hdr,  9, bold=True, color=C_HDR)
        tb(x, 1.22, 2.98, 4.25, body, 8, color=C_BODY)

    add_ai_badge(slide)
    return len(prs.slides) - 1


# ─────────────────────────────────────────────────────────────────────────────
prs = Presentation(SRC)
print(f'Opened v31: {len(prs.slides)} slides')

# ── 1. Fix S62 title ──────────────────────────────────────────────────────────
idx = find_idx(prs, "Supervisor's variable has a neural signal")
if idx is not None:
    sl = prs.slides[idx]
    for sh in sl.shapes:
        if sh.has_text_frame and "Supervisor's variable" in sh.text_frame.text:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if "Supervisor's variable" in run.text:
                        run.text = run.text.replace(
                            "Supervisor's variable has a neural signal",
                            "Neural signal: E07 × Cue Visible, E23 × Upcoming Choice")
            print(f'  S{idx+1}: title updated')
            break

# ── 2. Add AI-generated badge to all appendix slides (idx 68–85) ─────────────
APPENDIX_START = 68
for i in range(APPENDIX_START, len(prs.slides)):
    add_ai_badge(prs.slides[i])
print(f'  Added AI-generated badge to S{APPENDIX_START+1}–S{len(prs.slides)}')

# ── 3. Update A02 (GPV) — add GPV/R² note and S63 to APPEARS ON ──────────────
idx_a02 = find_idx(prs, 'A02 · Metric: GPV')
if idx_a02 is not None:
    sl = prs.slides[idx_a02]
    # Shape 7 = COMPUTATION body (idx 7), Shape 9 = APPEARS ON body (idx 9)
    for sh in sl.shapes:
        if sh.has_text_frame and 'APPEARS ON' in sh.text_frame.text:
            tf = sh.text_frame
            # Add S63 reference
            last_para = tf.paragraphs[-1]
            from pptx.oxml.ns import qn
            from lxml import etree
            new_p = etree.SubElement(tf._txBody, qn('a:p'))
            r_elm = etree.SubElement(new_p, qn('a:r'))
            rPr = etree.SubElement(r_elm, qn('a:rPr'),
                                    attrib={'lang': 'en-US', 'dirty': '0'})
            rPr.set('sz', '800')
            t = etree.SubElement(r_elm, qn('a:t'))
            t.text = 'S63 Bucket 1 (GPV/R² ≥ 10%)'
    for sh in sl.shapes:
        if sh.has_text_frame and 'COMPUTATION' in sh.text_frame.text:
            tf = sh.text_frame
            from pptx.oxml.ns import qn
            from lxml import etree
            new_p = etree.SubElement(tf._txBody, qn('a:p'))
            r_elm = etree.SubElement(new_p, qn('a:r'))
            rPr = etree.SubElement(r_elm, qn('a:rPr'),
                                    attrib={'lang': 'en-US', 'dirty': '0'})
            rPr.set('sz', '800')
            t = etree.SubElement(r_elm, qn('a:t'))
            t.text = 'GPV/R²: normalise by ensemble R² to compare across ensembles; threshold ≥ 0.10 used in bucket analysis (S63/S64).'
    print(f'  S{idx_a02+1}: A02 updated')

# ── 4. Update A09 (η²/Cohen's d) — add bucket usage ─────────────────────────
idx_a09 = find_idx(prs, "A09 · Metric: η²")
if idx_a09 is not None:
    sl = prs.slides[idx_a09]
    for sh in sl.shapes:
        if sh.has_text_frame and 'APPEARS ON' in sh.text_frame.text:
            tf = sh.text_frame
            from pptx.oxml.ns import qn
            from lxml import etree
            for line in ["S62 d ≥ 0.1 session filter",
                         "S64 η²×GPV/R² joint score (Bucket 2)"]:
                new_p = etree.SubElement(tf._txBody, qn('a:p'))
                r_elm = etree.SubElement(new_p, qn('a:r'))
                rPr = etree.SubElement(r_elm, qn('a:rPr'),
                                        attrib={'lang': 'en-US', 'dirty': '0'})
                rPr.set('sz', '800')
                t = etree.SubElement(r_elm, qn('a:t'))
                t.text = line
    print(f'  S{idx_a09+1}: A09 updated')

# ── 5. Add A18 slide: Bucket Analysis Methodology ────────────────────────────
a18_idx = add_appendix_slide(
    prs,
    tag='A18',
    title='A18 · Methodology: E07/E23 Bucket Analysis',
    col1_hdr='STEP 1 — SIGNAL FILTER',
    col1_body=(
        'For E07 (ensemble 6 × cue_visible) and E23 (ensemble 22 × upcoming_choice):\n\n'
        'Compute Cohen\'s d (max pairwise effect) between condition groups per session.\n\n'
        'Threshold: d ≥ 0.1\n'
        'Defines the analysis set — sessions where the (ensemble, feature) pair has a detectable neural signal.\n\n'
        'E07: 9/19 sessions above threshold\n'
        'E23: 19/23 sessions above threshold'
    ),
    col2_hdr='STEP 2 — BUCKET SPLIT',
    col2_body=(
        'For analysis-set sessions, compute GPV(feature)/R² — the fraction of ensemble predictive variance attributed to that feature.\n\n'
        'Bucket 1 (direct attribution):\n'
        'GPV(feature)/R² ≥ 0.10\n'
        'Model directly uses that feature.\n'
        'E07: 7/9. E23: 0/19.\n\n'
        'Bucket 2 (indirect):\n'
        'GPV(feature)/R² < 0.10\n'
        'Model does not directly attribute to this feature. Goes to Step 3.'
    ),
    col3_hdr='STEP 3 — CO-VARIATION',
    col3_body=(
        'For Bucket 2 sessions, find the feature g* that best explains the signal through joint co-variation + attribution:\n\n'
        'g* = argmax_g [ η²(g, condition) × GPV(g)/R² ]\n\n'
        'Plot:\n'
        'X = η²(g*, condition) — co-variation of g* with the condition label\n'
        'Y = GPV(g*)/R² — model attribution to g*\n\n'
        'Upper-right quadrant: the model uses a feature that co-varies with the condition → signal is captured, just through a different variable.'
    )
)
print(f'  Added A18 at idx {a18_idx}  (S{a18_idx+1})')

# ── Verify ────────────────────────────────────────────────────────────────────
print(f'\nFinal slide count: {len(prs.slides)}')
print('\nS62–S64:')
for frag in ['Neural signal: E07', 'Bucket 1: model', 'Bucket 2: a co']:
    i = find_idx(prs, frag)
    if i is not None:
        print(f'  S{i+1}: {prs.slides[i].shapes[0].text_frame.text[:60] if prs.slides[i].shapes else "?"}')
print(f'\nAppendix: S{APPENDIX_START+1}–S{len(prs.slides)} ({len(prs.slides)-APPENDIX_START} slides)')

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'\nSaved → {out}')
