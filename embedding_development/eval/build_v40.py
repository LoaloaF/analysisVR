#!/usr/bin/env python3
"""
build_v40.py  —  Replace S48 with updated cross-seed GPV figure (all 3 models).

Base: user's Desktop ultimate_presentation_v39.pptx
v39 (93) → v40 (93 slides, same count):
  S48: replace attribution_consistency_merged.png
       Right panel now shows MLP + TC-Cont + TC-Pred cross-seed GPV agreement
"""
import os, zipfile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.parts.presentation import PresentationPart


@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
mdir = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')

SRC  = '/mnt/c/Users/amits/Desktop/ultimate_presentation_v39.pptx'
OUTS = [
    os.path.join(root, 'outputs', 'ultimate_presentation_v40.pptx'),
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v40.pptx',
]
IMG  = os.path.join(mdir, 'attribution_consistency_merged.png')
TARGET_DPI = 200


def verify_figure(path, label):
    if not os.path.exists(path):
        print(f'  ✗ MISSING: {label}'); return False
    with Image.open(path) as im:
        w, h = im.size
    print(f'  ✓ {label}: {w}×{h}px = {w/TARGET_DPI:.2f}"×{h/TARGET_DPI:.2f}"')
    return True


def verify_zip(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        dupes = [n for n in set(names) if names.count(n) > 1]
        slides = [n for n in names if 'slides/slide' in n
                  and n.endswith('.xml') and 'rels' not in n]
    ok = len(dupes) == 0
    print(f'  {"✓" if ok else "✗ DUPES: "+str(dupes)} {len(slides)} slides, {len(dupes)} dupes')
    return ok


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


print('PRE-BUILD VERIFICATION')
if not verify_figure(IMG, 'attribution_consistency_merged'):
    raise SystemExit('Missing figure')
if not os.path.exists(SRC):
    raise SystemExit(f'Source not found: {SRC}')
print(f'  ✓ source: {os.path.basename(SRC)}\n')

prs = Presentation(SRC)
print(f'Opened: {len(prs.slides)} slides')

idx = find_idx(prs, 'Attribution Profile Consistency')
if idx is None:
    raise SystemExit('Could not find Attribution Profile Consistency slide')

sl = prs.slides[idx]
for sh in list(sl.shapes):
    if sh.shape_type == 13:
        sh._element.getparent().remove(sh._element)
        break
sl.shapes.add_picture(IMG,
    int(Inches(0.30)), int(Inches(0.62)),
    int(Inches(9.40)), int(Inches(4.16)))
print(f'  S{idx+1}: replaced attribution_consistency_merged.png')

print(f'\nFinal: {len(prs.slides)} slides')
for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved → {out}')

print('\nPOST-BUILD VERIFICATION')
for out in OUTS:
    if os.path.exists(out):
        verify_zip(out)
