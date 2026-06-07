#!/usr/bin/env python3
"""
build_v14.py  —  Apply supervisor feedback (v10 items) to v13 (68 slides).

Changes:
  1. S17 (idx 16): subtitle "MLP · TempConv" → clarify two separate model families
  2. S18 (idx 17): title too verbose ("LLM energy") → shorter direct title
  3. S22 (idx 21): replace figure with regenerated ablation_proof.png (frame_raw fixed)
  4. S36 (idx 35): section title reframed per supervisor suggestion
  5. S38 (idx 37): replace heatmap with gpv_task_ensembles.png (highlights task-sensitive ensembles)

Net: 68 slides (no count change — all in-place replacements)
"""
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
mdir = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')
cdir = os.path.join(root, 'outputs', 'cebra_comparison')
ml_naive_dir = os.path.join(root, 'outputs', 'mlps', 'ml_vs_naive')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v13.pptx')
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v14.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation_v14.pptx'),
]

SLIDE_W = Inches(10.0)
SLIDE_H = Inches(5.625)
DPI     = 200


def _fit_in_content_area(fw, fh, max_w=Inches(9.5), max_h=Inches(4.8),
                          top=Inches(0.75)):
    scale = min(max_w / Inches(fw), max_h / Inches(fh))
    w = int(Inches(fw) * scale)
    h = int(Inches(fh) * scale)
    l = int((SLIDE_W - w) // 2)
    t = int(top + (max_h - h) // 2)
    return l, t, w, h


def replace_figure(slide, img_path):
    pics = [s for s in slide.shapes if s.shape_type == 13]
    if not pics:
        print(f'  WARNING: no picture on slide')
        return
    target = max(pics, key=lambda s: s.width * s.height)
    l, t, w, h = int(target.left), int(target.top), int(target.width), int(target.height)
    target._element.getparent().remove(target._element)
    slide.shapes.add_picture(img_path, l, t, w, h)


def patch_text(slide, old_fragment, new_text):
    """Replace first run containing old_fragment with new_text."""
    for sh in slide.shapes:
        try:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if old_fragment in run.text:
                        run.text = run.text.replace(old_fragment, new_text)
                        return True
        except Exception:
            pass
    return False


def set_first_matching_run(slide, match_fn, new_text):
    """Set the first run whose text satisfies match_fn to new_text."""
    for sh in slide.shapes:
        try:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if match_fn(run.text):
                        run.text = new_text
                        return True
        except Exception:
            pass
    return False


# ── Open ──────────────────────────────────────────────────────────────────────
prs = Presentation(SRC)
print(f'Opened v13: {len(prs.slides)} slides')

# ── Fix 1: S17 (idx 16) — subtitle naming clarity ─────────────────────────────
# Runs are split: run[0]='MLP · '  run[1]='TempConv'
for sh in prs.slides[16].shapes:
    try:
        tf = sh.text_frame
        if tf.text.strip() == 'MLP · TempConv':
            paras = tf.paragraphs
            if len(paras[0].runs) >= 2:
                paras[0].runs[0].text = 'MLP (single time point)  ·  '
                paras[0].runs[1].text = 'TempConv-Cont / TempConv-Pred (temporal context)'
                print('  S17 subtitle: updated')
                break
    except Exception:
        pass
else:
    print('  S17 subtitle: NOT FOUND')

# ── Fix 2: S18 (idx 17) — shorten title ───────────────────────────────────────
ok = set_first_matching_run(
    prs.slides[17],
    lambda t: 'Attribution-Flagged' in t,
    'MLP Outperforms Linear Regression on Nonlinear Pairs'
)
print(f'  S18 title: {"updated" if ok else "NOT FOUND"}')

# ── Fix 3: S22 (idx 21) — replace ablation figure (frame_raw fixed) ───────────
ablation_img = os.path.join(ml_naive_dir, 'ablation_proof.png')
if os.path.exists(ablation_img):
    replace_figure(prs.slides[21], ablation_img)
    print(f'  S22 figure: replaced with fixed ablation_proof.png')
else:
    print(f'  S22 figure: SKIP ({ablation_img} not found)')

# ── Fix 4: S36 (idx 35) — reframe section title ───────────────────────────────
ok = set_first_matching_run(
    prs.slides[35],
    lambda t: t.strip() == 'Feature Attribution',
    'What Does the Best Model Actually Encode?'
)
print(f'  S36 title: {"updated" if ok else "NOT FOUND"}')

# ── Fix 5: S38 (idx 37) — replace heatmap with task-ensemble version ──────────
task_img = os.path.join(mdir, 'gpv_task_ensembles.png')
if os.path.exists(task_img):
    replace_figure(prs.slides[37], task_img)
    # Update title
    set_first_matching_run(
        prs.slides[37],
        lambda t: 'Global Permutation Variance' in t,
        'GPV by Feature Group × Ensemble: Task-Sensitive Ensembles Highlighted'
    )
    print(f'  S38 figure: replaced with gpv_task_ensembles.png')
else:
    print(f'  S38 figure: SKIP ({task_img} not found)')

# ── Save ──────────────────────────────────────────────────────────────────────
print(f'\nFinal slide count: {len(prs.slides)}')
for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved  → {out}')
