#!/usr/bin/env python3
"""
build_v36.py  —  Insert frequency trace example slide after S35.

v35 (88) → v36 (89 slides, +1):
  Insert S36: "TempConv-Pred tracks fast fluctuations — MLP cannot"
    Figure: freq_trace_example.png (S01 E03; noise R²: MLP=-0.10, TC-Pred=+0.26)
    Placed immediately after S35 (trend_noise_comparison bar chart)
  S36+ shifts by 1.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.parts.presentation import PresentationPart


@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
mdir = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v35.pptx')
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v36.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation_v36.pptx'),
]

C_BG    = RGBColor(0xFF, 0xFF, 0xFF)
C_TITLE = RGBColor(0x1A, 0x23, 0x3A)


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def move_slide(prs, from_idx, to_idx):
    xml = prs.slides._sldIdLst
    els = list(xml); el = els.pop(from_idx); els.insert(to_idx, el)
    for c in list(xml): xml.remove(c)
    for c in els:        xml.append(c)


prs = Presentation(SRC)
print(f'Opened v35: {len(prs.slides)} slides')

img_path = os.path.join(mdir, 'freq_trace_example.png')

slide = prs.slides.add_slide(prs.slide_layouts[6])
fill  = slide.background.fill; fill.solid(); fill.fore_color.rgb = C_BG

tb = slide.shapes.add_textbox(Inches(0.20), Inches(0.06), Inches(9.60), Inches(0.52))
p  = tb.text_frame.paragraphs[0]; r = p.add_run()
r.text = 'TempConv-Pred tracks fast fluctuations — MLP cannot'
r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = C_TITLE

if os.path.exists(img_path):
    slide.shapes.add_picture(img_path,
        int(Inches(0.25)), int(Inches(0.62)),
        int(Inches(9.50)), int(Inches(4.20)))
else:
    print(f'  WARNING: {img_path} not found')

new_idx = len(prs.slides) - 1

# Insert after S35 (TempConv-Pred bar chart)
after_idx = find_idx(prs, 'TempConv-Pred Captures')
if after_idx is None:
    after_idx = find_idx(prs, 'Fast Neural Fluctuations')
if after_idx is None:
    after_idx = find_idx(prs, 'Frequency Encoding')
    print(f'  Falling back to Frequency Encoding header at S{after_idx+1}')

insert_at = after_idx + 1
move_slide(prs, new_idx, insert_at)
print(f'  Inserted trace slide at S{insert_at+1} (after S{after_idx+1})')

print(f'\nFinal slide count: {len(prs.slides)}')
print('Context:')
for i in range(max(0, insert_at - 1), min(len(prs.slides), insert_at + 3)):
    sl = prs.slides[i]
    texts = [sh.text_frame.text.strip()[:70] for sh in sl.shapes if sh.has_text_frame]
    print(f'  S{i+1:02d}: {next((t for t in texts if t), "(empty)")}')

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'\nSaved → {out}')
