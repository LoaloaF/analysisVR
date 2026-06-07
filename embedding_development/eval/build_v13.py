#!/usr/bin/env python3
"""
build_v13.py  —  Add embedding consistency section to v12 (67 slides).

Changes from v12:
  1. S49 (idx 48): Replace old decodability-profile slide with
                   embedding_linear_map.png  (4-axis geometric consistency)
  2. NEW S50     : Insert cross_ensemble_prediction.png
                   (self vs cross-target prediction R²)

Net: 67 + 1 = 68 slides
"""
import os, shutil, zipfile, xml.etree.ElementTree as ET
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN

base  = os.path.dirname(os.path.abspath(__file__))
root  = os.path.join(base, '..')
mdir  = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')
cdir  = os.path.join(root, 'outputs', 'cebra_comparison')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v12.pptx')
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v13.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation_v13.pptx'),
]

SLIDE_W = Inches(10.0)
SLIDE_H = Inches(5.625)
DPI     = 200


def _img_size(img_path):
    img = Image.open(img_path)
    return img.width / DPI, img.height / DPI


def _fit_in_content_area(fw, fh,
                          max_w=Inches(9.5), max_h=Inches(4.8),
                          top=Inches(0.75)):
    scale = min(max_w / Inches(fw), max_h / Inches(fh))
    w = Inches(fw) * scale
    h = Inches(fh) * scale
    l = (SLIDE_W - w) // 2
    t = top + (max_h - h) // 2
    return int(l), int(t), int(w), int(h)


def replace_figure(slide, img_path):
    pics = [s for s in slide.shapes if s.shape_type == 13]
    if not pics:
        print(f'  WARNING: no picture shape found on slide')
        return
    target = max(pics, key=lambda s: s.width * s.height)
    l, t, w, h = int(target.left), int(target.top), int(target.width), int(target.height)
    sp = target._element
    sp.getparent().remove(sp)
    slide.shapes.add_picture(img_path, l, t, w, h)


def add_figure_slide(prs, img_path, title_text, subtitle_text=None):
    layout = prs.slide_layouts[6]
    slide  = prs.slides.add_slide(layout)
    from pptx.dml.color import RGBColor
    fill   = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

    tb = slide.shapes.add_textbox(Inches(0.3), Inches(0.08),
                                   Inches(9.4), Inches(0.5))
    tf = tb.text_frame; tf.word_wrap = False
    p  = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(16); p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    if subtitle_text:
        tb2 = slide.shapes.add_textbox(Inches(0.3), Inches(0.60),
                                        Inches(9.4), Inches(0.35))
        tf2 = tb2.text_frame; tf2.word_wrap = True
        p2  = tf2.paragraphs[0]
        p2.text = subtitle_text
        p2.font.size = Pt(10)
        p2.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)

    fw, fh = _img_size(img_path)
    l, t, w, h = _fit_in_content_area(fw, fh)
    slide.shapes.add_picture(img_path, l, t, w, h)
    return len(prs.slides) - 1


def move_slide(prs, from_idx, to_idx):
    xml_slides = prs.slides._sldIdLst
    children   = list(xml_slides)
    el         = children.pop(from_idx)
    children.insert(to_idx, el)
    for c in list(xml_slides):
        xml_slides.remove(c)
    for c in children:
        xml_slides.append(c)


def _clean_orphaned_parts(pptx_path):
    NS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    slide_rel = f'{{{NS}}}Relationship'
    with zipfile.ZipFile(pptx_path, 'r') as z:
        names     = set(z.namelist())
        rels_data = z.read('ppt/_rels/presentation.xml.rels')
    tree      = ET.fromstring(rels_data)
    referenced = set()
    for rel in tree.iter(slide_rel):
        t  = rel.get('Type', '')
        tg = rel.get('Target', '')
        if 'relationships/slide' in t:
            referenced.add('ppt/' + tg.lstrip('/'))
    slide_entries = [n for n in names
                     if 'ppt/slides/slide' in n
                     and 'slideLayout' not in n
                     and 'slideMaster' not in n]
    orphans = set()
    for name in slide_entries:
        bare = name.replace('/slides/_rels/', '/slides/').replace('.rels', '')
        if bare not in referenced:
            orphans.add(name)
            orphans.add(name.replace('/slides/', '/slides/_rels/').replace('.xml', '.xml.rels'))
    if not orphans:
        return
    tmp = pptx_path + '.tmp'
    with zipfile.ZipFile(pptx_path, 'r') as zin, \
         zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename not in orphans:
                zout.writestr(item, zin.read(item.filename))
    os.replace(tmp, pptx_path)


# ── Open ──────────────────────────────────────────────────────────────────────
prs = Presentation(SRC)
n_orig = len(prs.slides)
print(f'Opened v12: {n_orig} slides')

# ══════════════════════════════════════════════════════════════════════════════
# STEP 1 — Replace S49 (idx 48): old decodability slide → embedding_linear_map
# ══════════════════════════════════════════════════════════════════════════════
linmap_img = os.path.join(mdir, 'embedding_linear_map.png')
if os.path.exists(linmap_img):
    replace_figure(prs.slides[48], linmap_img)
    # Update title text box if present
    for shape in prs.slides[48].shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            tf = shape.text_frame
            for para in tf.paragraphs:
                for run in para.runs:
                    if 'Decodability' in run.text or 'Consistency' in run.text or 'consistency' in run.text:
                        run.text = 'Embedding Geometric Consistency: Linear Map R² Across Variation Axes'
            break
    print('  Replaced S49 with embedding_linear_map.png')
else:
    print(f'  SKIP: {linmap_img} not found')

# ══════════════════════════════════════════════════════════════════════════════
# STEP 2 — Add cross_ensemble_prediction.png as new slide (append, then move)
# ══════════════════════════════════════════════════════════════════════════════
pred_img = os.path.join(mdir, 'cross_ensemble_prediction.png')
if os.path.exists(pred_img):
    new_idx = add_figure_slide(
        prs, pred_img,
        'Embedding Consistency: Cross-Ensemble Neural Prediction',
        'Does an embedding trained on one neural ensemble predict other ensembles? '
        'Self vs cross-target R² and source×target heatmap (MLP, seed=42).'
    )
    # Move to S50 (idx 49), right after the linear map slide
    move_slide(prs, new_idx, 49)
    print(f'  Added cross_ensemble_prediction.png at S50 (idx 49)')
else:
    print(f'  SKIP: {pred_img} not found')

# ── Save ──────────────────────────────────────────────────────────────────────
print(f'\nFinal slide count: {len(prs.slides)}  (was {n_orig})')
for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved  → {out}')
