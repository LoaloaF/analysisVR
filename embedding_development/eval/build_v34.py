#!/usr/bin/env python3
"""
build_v34.py  —  Add E07 joint tuning slide after bucket analysis section.

v33 (87) → v34 (88 slides, +1):
  Insert S65: "E07: joint position × cue encoding — a challenge for marginal attribution"
    Figure: e07_joint_tuning.png
    Subtitle note: why marginal GPV misses it; grouped GPV as the solution
  S65+ (Summary onwards) shifts by 1.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.parts.presentation import PresentationPart


@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
adir = os.path.join(root, 'outputs', 'ablation_vs_attribution')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v33.pptx')
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v34.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation_v34.pptx'),
]

C_BG    = RGBColor(0xFF, 0xFF, 0xFF)
C_TITLE = RGBColor(0x1A, 0x23, 0x3A)
C_NOTE  = RGBColor(0x44, 0x44, 0x44)


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def move_slide(prs, from_idx, to_idx):
    xml = prs.slides._sldIdLst
    els = list(xml); el = els.pop(from_idx); els.insert(to_idx, el)
    for c in list(xml): xml.remove(c)
    for c in els:        xml.append(c)


prs = Presentation(SRC)
print(f'Opened v33: {len(prs.slides)} slides')

# Build new slide
img_path = os.path.join(adir, 'e07_joint_tuning.png')
slide = prs.slides.add_slide(prs.slide_layouts[6])
fill  = slide.background.fill; fill.solid(); fill.fore_color.rgb = C_BG

# Title
tb = slide.shapes.add_textbox(Inches(0.20), Inches(0.06), Inches(9.60), Inches(0.38))
p = tb.text_frame.paragraphs[0]; r = p.add_run()
r.text = 'E07: joint position × cue encoding'
r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = C_TITLE

# Note line below title
tb2 = slide.shapes.add_textbox(Inches(0.20), Inches(0.44), Inches(9.60), Inches(0.22))
p2 = tb2.text_frame.paragraphs[0]; r2 = p2.add_run()
r2.text = ('Marginal GPV permutes cue OR position separately — misses their interaction.  '
           'Grouped GPV (permute both together) would capture the joint contribution.')
r2.font.size = Pt(9); r2.font.italic = True; r2.font.color.rgb = C_NOTE

# Figure
if os.path.exists(img_path):
    slide.shapes.add_picture(img_path,
        int(Inches(0.25)), int(Inches(0.68)),
        int(Inches(9.50)), int(Inches(4.20)))
else:
    print(f'  WARNING: {img_path} not found')

new_idx = len(prs.slides) - 1
print(f'  Added joint tuning slide at idx {new_idx}')

# Move to immediately after the last bucket slide (Bucket 2)
target_frag = 'Bucket 2: a co-varying'
after_idx   = find_idx(prs, target_frag)
if after_idx is None:
    after_idx = find_idx(prs, 'Case Studies')
    print(f'  WARNING: bucket 2 slide not found, inserting after Case Studies')

insert_at = after_idx + 1
move_slide(prs, new_idx, insert_at)
print(f'  Moved to idx {insert_at} (after S{after_idx+1})')

# Verify
print(f'\nFinal slide count: {len(prs.slides)}')
print('Slides around insertion:')
for i in range(max(0, insert_at - 1), min(len(prs.slides), insert_at + 3)):
    sl = prs.slides[i]
    texts = [sh.text_frame.text.strip()[:70] for sh in sl.shapes if sh.has_text_frame]
    print(f'  S{i+1:02d}: {next((t for t in texts if t), "(empty)")}')

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'\nSaved → {out}')
