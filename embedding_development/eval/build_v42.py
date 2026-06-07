#!/usr/bin/env python3
"""
build_v42.py  —  Fix S36 legend overlap, S68 title, S69 → joint GPV, delete S70.

Base: user's Desktop ultimate_presentation_v41.pptx
v41 (93) → v42 (92 slides, -1):
  S36: replace freq_trace_example.png (opaque legend fix)
  S68: update title → "E07 — cue-specific position tuning (top 4 sessions)"
  S69: replace empty averaged-curves with e07_joint_gpv.png
       title → "Joint attribution: cue+position interaction ≈ 0 (additive)"
  S70: delete (redundant with bucket-2 case studies on S65/S66)
"""
import os, zipfile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.parts.presentation import PresentationPart


@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
mdir = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')

SRC  = '/mnt/c/Users/amits/Desktop/ultimate_presentation_v41.pptx'
OUTS = [
    os.path.join(root, 'outputs', 'ultimate_presentation_v42.pptx'),
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v42.pptx',
]

IMG_FREQ  = os.path.join(mdir, 'freq_trace_example.png')
IMG_GJPV  = '/mnt/c/Users/amits/Desktop/e07_joint_gpv.png'
TARGET_DPI = 200

TITLE_S68 = 'E07 — cue-specific position tuning (top 4 sessions)'
TITLE_S69 = 'Joint attribution: cue+position interaction ≈0 (additive)'


def verify_figure(path, label):
    if not os.path.exists(path):
        print(f'  ✗ MISSING: {label}'); return False
    with Image.open(path) as im:
        w, h = im.size
    print(f'  ✓ {label}: {w}×{h}px = {w/TARGET_DPI:.2f}"×{h/TARGET_DPI:.2f}"')
    return True


def verify_zip(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        dupes = [n for n in set(names) if names.count(n) > 1]
        slides = [n for n in names if 'slides/slide' in n
                  and n.endswith('.xml') and 'rels' not in n]
    ok = len(dupes) == 0
    print(f'  {"✓" if ok else "✗ DUPES: "+str(dupes)} {len(slides)} slides, {len(dupes)} dupes')
    return ok


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def replace_image(slide, new_path, left, top, w, h):
    for sh in list(slide.shapes):
        if sh.shape_type == 13:
            sh._element.getparent().remove(sh._element)
            break
    slide.shapes.add_picture(new_path,
        int(Inches(left)), int(Inches(top)),
        int(Inches(w)),    int(Inches(h)))


def patch_title(slide, new_text):
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            if para.runs:
                para.runs[0].text = new_text
                for r in list(para.runs)[1:]:
                    r.text = ''
                return True
    return False


def delete_slide(prs, idx):
    xml = prs.slides._sldIdLst
    els = list(xml)
    xml.remove(els[idx])


print('PRE-BUILD VERIFICATION')
ok  = verify_figure(IMG_FREQ, 'freq_trace_example')
ok &= verify_figure(IMG_GJPV, 'e07_joint_gpv')
ok &= len(TITLE_S68) <= 75
ok &= len(TITLE_S69) <= 75
print(f'  ✓ S68 title: {len(TITLE_S68)} chars')
print(f'  ✓ S69 title: {len(TITLE_S69)} chars')
if not os.path.exists(SRC):
    raise SystemExit(f'Source not found: {SRC}')
print(f'  ✓ source: {os.path.basename(SRC)}\n')
if not ok:
    raise SystemExit('Verification failed')

prs = Presentation(SRC)
print(f'Opened: {len(prs.slides)} slides')

# ── S36: replace freq trace image (the "Example Plots" slide after the MLP bar chart)
# S35=MLP Fails bar chart, S36=freq trace example — find by large image in the range
idx36 = None
for probe in range(32, 40):
    sl = prs.slides[probe]
    for sh in sl.shapes:
        if sh.shape_type == 13 and len(sh.image.blob) > 280000:
            idx36 = probe; break
    if idx36 is not None: break
if idx36 is None:
    raise SystemExit('Could not find freq trace slide (S36 area)')
replace_image(prs.slides[idx36], IMG_FREQ, 0.25, 0.62, 9.50, 4.20)
print(f'  S{idx36+1}: freq_trace_example replaced')

# ── S68: update title ─────────────────────────────────────────────────────────
idx68 = find_idx(prs, 'joint position')
if idx68 is None:
    idx68 = find_idx(prs, 'E07: joint')
if idx68 is None:
    raise SystemExit('Could not find S68 (joint position × cue encoding)')
patched = patch_title(prs.slides[idx68], TITLE_S68)
print(f'  S{idx68+1}: title patched={patched} → "{TITLE_S68}"')

# ── S69: replace empty averaged curves with joint GPV figure ─────────────────
idx69 = find_idx(prs, 'position tuning averaged')
if idx69 is None:
    idx69 = find_idx(prs, 'averaged across sessions')
if idx69 is None:
    raise SystemExit('Could not find S69 (averaged tuning curves)')
replace_image(prs.slides[idx69], IMG_GJPV, 0.25, 0.62, 9.50, 4.20)
patched = patch_title(prs.slides[idx69], TITLE_S69)
print(f'  S{idx69+1}: joint GPV image inserted, title patched={patched}')

# ── S70: delete "Signal routes through speed" ─────────────────────────────────
idx70 = find_idx(prs, 'Signal routes through speed')
if idx70 is None:
    idx70 = find_idx(prs, 'routes through speed')
if idx70 is None:
    print('  WARNING: could not find S70 "Signal routes through speed" — skipping delete')
else:
    delete_slide(prs, idx70)
    print(f'  S{idx70+1}: deleted')

print(f'\nFinal: {len(prs.slides)} slides')

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved → {out}')

print('\nPOST-BUILD VERIFICATION')
for out in OUTS:
    if os.path.exists(out):
        verify_zip(out)
