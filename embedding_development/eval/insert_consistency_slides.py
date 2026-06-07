#!/usr/bin/env python3
"""
insert_consistency_slides.py

Insert two embedding consistency comparison slides after S40 (TempConv-Pred R² bar,
idx=39) in Neural_Task_Representation_cleanup.pptx.

New slides:
  S41: "Embedding Consistency: All Models"    — heatmaps + violin
  S42: "Embedding Consistency: Per-Ensemble"  — scatter MLP vs TempConv
"""
import os
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from copy import deepcopy

PPTX_PATH = "/mnt/c/Users/amits/Downloads/Neural_Task_Representation_cleanup.pptx"
base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, "..")
cmp_dir = os.path.join(root, "outputs", "cebra_comparison")

IMG_L, IMG_T, IMG_W, IMG_H = Inches(0.25), Inches(1.05), Inches(9.5), Inches(5.5)

INSERT_AFTER = 39   # 0-based idx of S40 (TempConv-Pred R² bar)

new_slides = [
    ("Embedding Consistency: All Models",
     os.path.join(cmp_dir, "embedding_consistency_comparison.png")),
    ("Embedding Consistency: Per-Ensemble Scatter",
     os.path.join(cmp_dir, "embedding_consistency_scatter.png")),
]

prs = Presentation(PPTX_PATH)
layout = prs.slide_layouts[3]   # TITLE_ONLY


def _move_slide(prs, from_idx, to_idx):
    """Move slide from from_idx to to_idx (both 0-based) by reordering the XML."""
    xml_slides = prs.slides._sldIdLst
    slides_list = list(xml_slides)
    el = slides_list[from_idx]
    xml_slides.remove(el)
    xml_slides.insert(to_idx, el)


for i, (title_text, img_path) in enumerate(new_slides):
    if not os.path.exists(img_path):
        print(f"SKIP (missing): {img_path}")
        continue

    # Add slide at end
    slide = prs.slides.add_slide(layout)
    end_idx = len(prs.slides) - 1

    # Set title
    title_ph = slide.shapes.title
    if title_ph is not None:
        title_ph.text = title_text
        for para in title_ph.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(20)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)

    # Add image
    slide.shapes.add_picture(img_path, IMG_L, IMG_T, width=IMG_W, height=IMG_H)

    # Move to INSERT_AFTER + 1 + i
    target = INSERT_AFTER + 1 + i
    _move_slide(prs, end_idx, target)
    print(f"Inserted S{target+1} (idx={target}): '{title_text}'")

prs.save(PPTX_PATH)
print(f"\nSaved → {PPTX_PATH}  ({len(prs.slides)} slides total)")
