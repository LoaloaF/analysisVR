#!/usr/bin/env python3
"""
build_v41.py  —  Replace S67 with averaged E07 position × cue tuning curve.

Base: user's Desktop ultimate_presentation_v40.pptx
v40 (93) → v41 (93 slides, same count):
  S67: replace e07_assembly_heatmap.png with e07_position_cue_tuning.png
       Update title: "E07 × Cue 2: model captures the spatial pattern"
                  → "E07 — position tuning averaged across sessions"
"""
import os, zipfile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.parts.presentation import PresentationPart


@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
adir = os.path.join(root, 'outputs', 'ablation_vs_attribution')

SRC  = '/mnt/c/Users/amits/Desktop/ultimate_presentation_v40.pptx'
OUTS = [
    os.path.join(root, 'outputs', 'ultimate_presentation_v41.pptx'),
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v41.pptx',
]
IMG   = os.path.join(adir, 'e07_position_cue_tuning.png')
TITLE = 'E07 — position tuning averaged across sessions'
TARGET_DPI = 200


def verify_figure(path, label):
    if not os.path.exists(path):
        print(f'  ✗ MISSING: {label}'); return False
    with Image.open(path) as im:
        w, h = im.size
    print(f'  ✓ {label}: {w}×{h}px = {w/TARGET_DPI:.2f}"×{h/TARGET_DPI:.2f}"')
    return True


def verify_zip(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        dupes = [n for n in set(names) if names.count(n) > 1]
        slides = [n for n in names if 'slides/slide' in n
                  and n.endswith('.xml') and 'rels' not in n]
    ok = len(dupes) == 0
    print(f'  {"✓" if ok else "✗ DUPES: "+str(dupes)} {len(slides)} slides, {len(dupes)} dupes')
    return ok


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def patch_title(slide, new_text):
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            if para.runs:
                para.runs[0].text = new_text
                for r in list(para.runs)[1:]:
                    r.text = ''
                return True
    return False


print('PRE-BUILD VERIFICATION')
ok = verify_figure(IMG, 'e07_position_cue_tuning')
ok &= len(TITLE) <= 75
print(f'  ✓ title: {len(TITLE)} chars')
if not os.path.exists(SRC):
    raise SystemExit(f'Source not found: {SRC}')
print(f'  ✓ source: {os.path.basename(SRC)}\n')
if not ok:
    raise SystemExit('Verification failed')

prs = Presentation(SRC)
print(f'Opened: {len(prs.slides)} slides')

idx = find_idx(prs, 'model captures the spatial pattern')
if idx is None:
    idx = find_idx(prs, 'Cue 2')
if idx is None:
    raise SystemExit('Could not find S67 target slide')

sl = prs.slides[idx]
for sh in list(sl.shapes):
    if sh.shape_type == 13:
        sh._element.getparent().remove(sh._element)
        break
sl.shapes.add_picture(IMG,
    int(Inches(0.25)), int(Inches(0.62)),
    int(Inches(9.50)), int(Inches(4.20)))
patched = patch_title(sl, TITLE)
print(f'  S{idx+1}: figure replaced, title patched={patched}')

print(f'\nFinal: {len(prs.slides)} slides')
for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved → {out}')

print('\nPOST-BUILD VERIFICATION')
for out in OUTS:
    if os.path.exists(out):
        verify_zip(out)
