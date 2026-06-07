#!/usr/bin/env python3
"""
build_v20.py  —  Drop S47 (probe-prediction consistency heatmaps) from v19.

S47 showed cross-seed Pearson r of Ridge probe *outputs* — not CEBRA's actual
consistency metric (linear map R² between embedding spaces).  S49 (Embedding
Geometric Consistency: Linear Map R²) already carries the correct story, so
S47 is redundant and misleading.

Change: 85 → 84 slides.
  Removed: S47 "Embedding Consistency Across Models: MLP, TempConv-Cont, TempConv-Pred"
  Old S48 → S47, Old S49 → S48, etc.
"""
import os
from pptx import Presentation

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v19.pptx')
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v20.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation_v20.pptx'),
]

DROP_IDX = 46   # S47 — "Embedding Consistency Across Models" (probe-prediction Pearson r)


def remove_slide(prs, idx):
    slide      = prs.slides[idx]
    slide_part = slide.part
    prs_part   = prs.slides.part

    rId = None
    for rel_id, rel in prs_part.rels.items():
        try:
            if rel.target_part is slide_part:
                rId = rel_id
                break
        except Exception:
            pass
    if rId is not None:
        prs_part.rels.pop(rId)

    xml_slides = prs.slides._sldIdLst
    el = list(xml_slides)[idx]
    xml_slides.remove(el)


prs = Presentation(SRC)
print(f'Opened v19: {len(prs.slides)} slides')

# Confirm target before deleting
target = prs.slides[DROP_IDX]
texts  = [sh.text_frame.text.strip() for sh in target.shapes if sh.has_text_frame]
print(f'  Removing S{DROP_IDX+1} (idx {DROP_IDX}): {" | ".join(t for t in texts if t)[:120]}')

remove_slide(prs, DROP_IDX)
print(f'Final slide count: {len(prs.slides)}')

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved  → {out}')
