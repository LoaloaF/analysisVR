#!/usr/bin/env python3
"""
build_v46.py  —  Insert attribution evolution lineplot slide after S46.

Base: outputs/ultimate_presentation_v45.pptx  (93 slides)
v45 (93) → v46 (94 slides):
  New S47: evolution_lineplot_ig.png
           "Attribution evolves session-to-session: ensemble-specific fingerprints"
  Old S47–S93 shift to S48–S94.
"""
import os, zipfile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.parts.presentation import PresentationPart

@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v45.pptx')
OUTS = [
    os.path.join(root, 'outputs', 'ultimate_presentation_v46.pptx'),
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v46.pptx',
]

IMG  = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed',
                    'evolution_lineplot_ig.png')
TITLE = 'Attribution evolves session-to-session: ensemble-specific fingerprints'
TARGET_DPI = 200


def verify_figure(path, label):
    if not os.path.exists(path):
        print(f'  x MISSING: {label}'); return False
    with Image.open(path) as im:
        w, h = im.size
    print(f'  ok {label}: {w}x{h}px = {w/TARGET_DPI:.2f}"x{h/TARGET_DPI:.2f}"')
    return True


def verify_zip(path):
    with zipfile.ZipFile(path) as z:
        names  = z.namelist()
        dupes  = [n for n in set(names) if names.count(n) > 1]
        slides = [n for n in names if 'slides/slide' in n
                  and n.endswith('.xml') and 'rels' not in n]
    ok = len(dupes) == 0
    mark = 'ok' if ok else 'DUPES: ' + str(dupes)
    print(f'  {mark}  {len(slides)} slides, {len(dupes)} dupes')
    return ok


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def move_slide(prs, from_idx, to_idx):
    xml = prs.slides._sldIdLst
    els = list(xml)
    el  = els.pop(from_idx)
    els.insert(to_idx, el)
    for c in list(xml): xml.remove(c)
    for c in els:       xml.append(c)


def add_figure_slide(prs, after_idx, img_path, title_text):
    """Add a new slide with a bold title and a single figure image."""
    # Borrow layout from an adjacent content slide
    layout = prs.slides[after_idx].slide_layout
    new_slide = prs.slides.add_slide(layout)

    # Clear any placeholders from the layout
    for ph in list(new_slide.placeholders):
        sp = ph._element
        sp.getparent().remove(sp)

    slide_w = prs.slide_width
    slide_h = prs.slide_height

    # Title text box: full width, top strip
    title_h = Inches(0.52)
    txBox = new_slide.shapes.add_textbox(
        Inches(0.20), Inches(0.06), slide_w - Inches(0.40), title_h)
    tf = txBox.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = title_text
    run.font.bold   = True
    run.font.size   = Pt(18)
    run.font.color.rgb = RGBColor(0x1a, 0x1a, 0x1a)
    run.font.name   = 'Arial'

    # Figure image
    with Image.open(img_path) as im:
        img_w_px, img_h_px = im.size
    img_w_in = img_w_px / TARGET_DPI
    img_h_in = img_h_px / TARGET_DPI

    left_in = (10.0 - img_w_in) / 2   # centre horizontally
    top_in  = 0.62
    new_slide.shapes.add_picture(
        img_path,
        int(Inches(left_in)), int(Inches(top_in)),
        int(Inches(img_w_in)), int(Inches(img_h_in)))

    return len(prs.slides) - 1   # index of newly added slide (currently at end)


# ── Pre-build verification ─────────────────────────────────────────────────
print('PRE-BUILD VERIFICATION')
ok = verify_figure(IMG, 'evolution_lineplot_ig')
assert len(TITLE) <= 80, f'Title too long: {len(TITLE)}'
print(f'  ok title: {len(TITLE)} chars')
if not os.path.exists(SRC):
    raise SystemExit(f'Source not found: {SRC}')
print(f'  ok source: {os.path.basename(SRC)}\n')
if not ok:
    raise SystemExit('Verification failed')

prs = Presentation(SRC)
print(f'Opened: {len(prs.slides)} slides')

# Find S46 by title fragment
idx46 = find_idx(prs, 'Attribution')
# S46 is the last attribution slide — search more specifically
for probe in ['Naive Effect Size', 'ml_vs_naive', 'captures novel', 'Novel Multi']:
    idx46 = find_idx(prs, probe)
    if idx46 is not None:
        break
if idx46 is None:
    raise SystemExit('Could not find S46 (Attribution ≠ Naive slide)')
print(f'  Found S46 at index {idx46} (S{idx46+1})')

# Add new slide (appended at end), then move into position
new_idx = add_figure_slide(prs, idx46, IMG, TITLE)
target  = idx46 + 1
move_slide(prs, new_idx, target)
print(f'  New evolution slide inserted at S{target+1}')

# Context check
print('\nAttribution section context:')
for i in range(max(0, idx46 - 1), min(len(prs.slides), idx46 + 5)):
    sl    = prs.slides[i]
    texts = [sh.text_frame.text.strip()[:60] for sh in sl.shapes
             if sh.has_text_frame and sh.text_frame.text.strip()]
    imgs  = any(sh.shape_type == 13 for sh in sl.shapes)
    print(f'  S{i+1:02d}: {texts[0] if texts else "(empty)"}{"  [IMG]" if imgs else ""}')

print(f'\nFinal: {len(prs.slides)} slides')
for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved -> {out}')

print('\nPOST-BUILD VERIFICATION')
for out in OUTS:
    if os.path.exists(out):
        verify_zip(out)
