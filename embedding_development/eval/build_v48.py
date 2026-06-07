#!/usr/bin/env python3
"""
build_v48.py — Update figures for S47, S51, S56, S70.

Base: outputs/ultimate_presentation_v47.pptx  (94 slides)
v47 (94) → v48 (94 slides):
  S47  evolution_lineplot_ig.png     — 6 panels, no legend, feature in title, smaller dots
  S51  cross_ensemble_prediction.png — violin-only (removed heatmap panel)
  S56  head_angle_scatter_2x2.png    — shortened axis labels
  S70  e07_cue_zone_joint_gpv.png    — new story: pos tuning cue-diff vs joint GPV
"""
import os, zipfile
from PIL import Image
from pptx import Presentation
from pptx.parts.presentation import PresentationPart

@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v47.pptx')
OUTS = [
    os.path.join(root, 'outputs', 'ultimate_presentation_v48.pptx'),
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v48.pptx',
]

TARGET_DPI = 200
mdir = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')
adir = os.path.join(root, 'outputs', 'ablation_vs_attribution')

UPDATES = [
    ('Attribution evolves session',          os.path.join(mdir, 'evolution_lineplot_ig.png'),       'S47'),
    ('Cross-Ensemble Prediction',            os.path.join(mdir, 'cross_ensemble_prediction.png'),   'S51'),
    ('Top Pairs: Head Angle',                os.path.join(mdir, 'head_angle_scatter_2x2.png'),      'S56'),
    ('Joint GPV in the cue zone',            os.path.join(adir, 'e07_cue_zone_joint_gpv.png'),      'S70'),
]


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def replace_figure(slide, new_img_path, prs, label):
    pic_shape = None
    max_area  = 0
    for shape in slide.shapes:
        if shape.shape_type == 13:
            area = shape.width * shape.height
            if area > max_area:
                max_area  = area
                pic_shape = shape
    if pic_shape is None:
        print(f'  WARNING: no picture found on {label}')
        return False
    old_top = pic_shape.top
    with Image.open(new_img_path) as im:
        iw, ih = im.size
    img_w_emu = int(iw * 914400 / TARGET_DPI)
    img_h_emu = int(ih * 914400 / TARGET_DPI)
    left_emu  = (prs.slide_width - img_w_emu) // 2
    pic_shape._element.getparent().remove(pic_shape._element)
    slide.shapes.add_picture(new_img_path, left_emu, old_top, img_w_emu, img_h_emu)
    print(f'      replaced  {iw}×{ih}px  ({iw/TARGET_DPI:.1f}"×{ih/TARGET_DPI:.1f}")')
    return True


def verify_zip(path):
    with zipfile.ZipFile(path) as z:
        names  = z.namelist()
        dupes  = [n for n in set(names) if names.count(n) > 1]
        slides = [n for n in names if 'slides/slide' in n
                  and n.endswith('.xml') and 'rels' not in n]
    ok = len(dupes) == 0
    print(f'  {"ok" if ok else "DUPES: "+str(dupes)}  {len(slides)} slides  {len(dupes)} dupes')
    return ok


print('PRE-BUILD CHECKS')
all_ok = True
for _, img_path, label in UPDATES:
    if os.path.exists(img_path):
        with Image.open(img_path) as im:
            print(f'  ok {label}: {os.path.basename(img_path)}  {im.size[0]}x{im.size[1]}px')
    else:
        print(f'  MISSING {label}: {img_path}')
        all_ok = False
if not all_ok:
    raise SystemExit('Missing figures — aborting')
if not os.path.exists(SRC):
    raise SystemExit(f'Source not found: {SRC}')
print(f'  ok source: {os.path.basename(SRC)}\n')

prs = Presentation(SRC)
print(f'Opened {os.path.basename(SRC)}: {len(prs.slides)} slides\n')

print('REPLACING FIGURES')
replaced = 0
for fragment, img_path, label in UPDATES:
    idx = find_idx(prs, fragment)
    if idx is None:
        print(f'  WARNING: could not find slide for "{fragment}" ({label})')
        continue
    print(f'  {label}  idx={idx}  "{fragment}"')
    if replace_figure(prs.slides[idx], img_path, prs, label):
        replaced += 1

print(f'\nReplaced {replaced}/{len(UPDATES)} figures')
print(f'Final: {len(prs.slides)} slides')
for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved -> {out}')

print('\nPOST-BUILD VERIFICATION')
for out in OUTS:
    if os.path.exists(out):
        verify_zip(out)
