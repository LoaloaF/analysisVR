#!/usr/bin/env python3
"""
build_v6.py  —  Build v6 from v5:
  1. Append new position_collinearity.png figure slide
  2. Move Parts 2/3/4 (nonlinearity justification) to after the Linear section
  3. Insert position_collinearity slide after position_tuning in case studies

v5 slide indices (0-based, 46 slides):
  0-10   Intro
  11,12  Linear
  13-15  MLP
  16-19  TempConv + R²
  20-22  Example predictions
  23-34  Attribution (23-31 = main, 32-34 = Parts 2/3/4)
  35,36  Consistency
  37     Frequency
  38-43  Case studies (head angle: 38-42, position: 43)
  44,45  E07/E23 case study

New v6 order (47 slides):
  0-10   Intro
  11,12  Linear
  32,33,34  Nonlinearity justification (Parts 2/3/4 — moved here)
  13-15  MLP
  16-19  TempConv + R²
  20-22  Example predictions
  23-31  Attribution (Part 1 + all attribution slides, without 32-34)
  35,36  Consistency
  37     Frequency
  38-43  Case studies head angle + position_tuning
  46     position_collinearity (new, appended as index 46)
  44,45  E07/E23 case study
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

SRC  = '/mnt/c/Users/amits/Desktop/ultimate_presentation_v5.pptx'
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v6.pptx',
    os.path.join(os.path.dirname(os.path.abspath(__file__)), '..',
                 'outputs', 'ultimate_presentation_v6.pptx'),
]

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
NEW_FIG_PATH = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed',
                            'position_collinearity.png')

SLIDE_W = Inches(10.0)
SLIDE_H = Inches(5.625)
TITLE_H = Inches(0.55)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
DARK    = RGBColor(0x1A, 0x23, 0x3A)
DPI     = 200


def add_figure_slide(prs, img_path, title_text):
    """Append a white-background figure slide matching v4 style."""
    layout = prs.slide_layouts[6]   # blank
    slide  = prs.slides.add_slide(layout)

    # White background
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE

    # Title bar
    tb = slide.shapes.add_textbox(
        Inches(0.15), Inches(0.05),
        SLIDE_W - Inches(0.3), TITLE_H)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title_text
    r.font.size = Pt(17)
    r.font.bold = True
    r.font.color.rgb = DARK

    # Figure image (centred in content area)
    img      = Image.open(img_path)
    fig_w    = img.width  / DPI
    fig_h    = img.height / DPI
    cont_top = TITLE_H + Inches(0.05)
    cont_h   = (SLIDE_H - cont_top - Inches(0.08)) / 914400
    cont_w   = (SLIDE_W - Inches(0.10)) / 914400
    scale    = min(cont_w / fig_w, cont_h / fig_h)
    pw       = Inches(fig_w * scale)
    ph       = Inches(fig_h * scale)
    left     = (SLIDE_W - pw) // 2
    top      = cont_top + Emu(int((cont_h - fig_h * scale) / 2 * 914400))
    slide.shapes.add_picture(img_path, left, top, pw, ph)

    # Filename caption
    cap = slide.shapes.add_textbox(
        SLIDE_W - Inches(4.0), SLIDE_H - Inches(0.22),
        Inches(3.9), Inches(0.20))
    p2 = cap.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    r2.text = os.path.basename(img_path)
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    return len(prs.slides) - 1


# ── Load v5, append new slide ─────────────────────────────────────────────────
prs = Presentation(SRC)
assert len(prs.slides) == 46, f"Expected 46 slides, got {len(prs.slides)}"

new_idx = add_figure_slide(
    prs, NEW_FIG_PATH,
    'Position Is Captured via Collinear Features — Task Events Are Position-Triggered')
print(f"Appended position_collinearity slide at index {new_idx} ({len(prs.slides)} total)")

# ── Desired order (all 0-based, 47 slides) ────────────────────────────────────
NEW_ORDER = [
    # Intro
    0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10,
    # Linear
    11, 12,
    # Nonlinearity justification (Parts 2/3/4, moved from attribution)
    32, 33, 34,
    # MLP
    13, 14, 15,
    # TempConv + R² comparison
    16, 17, 18, 19,
    # Example predictions
    20, 21, 22,
    # Attribution: intro + MLP + cross-model + Part 1 validation
    23, 24, 25, 26, 27, 28, 29, 30, 31,
    # Embedding consistency
    35, 36,
    # Frequency
    37,
    # Case studies: head angle, position tuning, collinearity
    38, 39, 40, 41, 42, 43, 46,
    # E07/E23
    44, 45,
]

assert len(NEW_ORDER) == 47, f"Expected 47, got {len(NEW_ORDER)}"
assert len(set(NEW_ORDER)) == 47, "Duplicate indices"

xml_slides = prs.slides._sldIdLst
all_els    = list(xml_slides)
for el in list(xml_slides):
    xml_slides.remove(el)
for idx in NEW_ORDER:
    xml_slides.append(all_els[idx])

print(f"Final slide count: {len(prs.slides)}")

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f"Saved → {out}")
print("Done.")
