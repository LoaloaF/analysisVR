#!/usr/bin/env python3
"""
convert_to_white_theme.py

Opens ultimate_presentation.pptx (preserving all user edits), converts the 22
figure slides (dark navy #1A233A background) to a white theme that matches the
story slides, and saves as ultimate_presentation_v2.pptx.

What changes on figure slides:
  - Background: dark navy → white (explicit solid #FFFFFF)
  - Title text:  white #FFFFFF → dark navy #1A233A  (bold preserved)
  - Caption/filename text: #888888 gray → #555555 darker gray

Story slides (no explicit bg fill) are left completely untouched.
"""

import os, sys
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL

# ── Paths ──────────────────────────────────────────────────────────────────────
SRC_PPTX = '/mnt/c/Users/amits/Desktop/ultimate_presentation.pptx'
OUT_PATHS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v2.pptx',
    os.path.join(os.path.dirname(os.path.abspath(__file__)), '..',
                 'outputs', 'ultimate_presentation_v2.pptx'),
]

# ── Colors ─────────────────────────────────────────────────────────────────────
DARK_NAVY   = RGBColor(0x1A, 0x23, 0x3A)   # original bg / new title color
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
WHITE_BG    = RGBColor(0xFF, 0xFF, 0xFF)
TITLE_COLOR = RGBColor(0x1A, 0x23, 0x3A)   # dark navy text on white
CAPTION_CLR = RGBColor(0x55, 0x55, 0x55)   # visible-but-subtle filename caption


def is_figure_slide(slide):
    """True if slide has our explicit dark navy background."""
    fill = slide.background.fill
    if fill.type == MSO_FILL.SOLID:
        try:
            return fill.fore_color.rgb == DARK_NAVY
        except Exception:
            pass
    return False


def convert_slide(slide, slide_num):
    # 1. Flip background to white
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = WHITE_BG

    # 2. Update every text run: white → dark navy; light gray → caption gray
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                try:
                    old = run.font.color.rgb
                except Exception:
                    continue  # no explicit color set — leave as-is

                r, g, b = old[0], old[1], old[2]

                # White / near-white → dark navy title color
                if r > 230 and g > 230 and b > 230:
                    run.font.color.rgb = TITLE_COLOR

                # Light gray caption (#888888 ±30) → darker gray
                elif 100 < r < 170 and abs(r - g) < 25 and abs(g - b) < 25:
                    run.font.color.rgb = CAPTION_CLR

    print(f"  Converted slide {slide_num} to white theme")


# ── Main ───────────────────────────────────────────────────────────────────────
print(f"Opening: {SRC_PPTX}")
prs = Presentation(SRC_PPTX)
print(f"  {len(prs.slides)} slides total")

n_converted = 0
for i, slide in enumerate(prs.slides):
    if is_figure_slide(slide):
        convert_slide(slide, i + 1)
        n_converted += 1

print(f"\n  {n_converted} figure slides converted; "
      f"{len(prs.slides) - n_converted} story slides left untouched")

for out_path in OUT_PATHS:
    os.makedirs(os.path.dirname(out_path) or '.', exist_ok=True)
    prs.save(out_path)
    print(f"Saved ({len(prs.slides)} slides) → {out_path}")

print("Done.")
