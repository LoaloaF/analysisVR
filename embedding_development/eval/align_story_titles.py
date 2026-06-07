#!/usr/bin/env python3
"""
align_story_titles.py

On v2, repositions and resizes the title placeholder on every story slide
(white bg / inherited fill) to match the figure-slide title:
  left=0.15"  top=0.05"  width=9.70"  height=0.55"  font=17pt bold

Story slides that have no title placeholder (slides 2, 9) are left untouched.
The title slide's ctrTitle is skipped (different layout purpose).

Saves as ultimate_presentation_v3.pptx.
"""

import os
from copy import deepcopy
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.oxml.ns import qn, nsmap
from pptx.enum.dml import MSO_FILL

# ── Target geometry matching figure slide title textbox ────────────────────────
L = int(Inches(0.15))    # 137 160 EMU
T = int(Inches(0.05))    #  45 720 EMU
W = int(Inches(9.70))    # 8 869 680 EMU
H = int(Inches(0.55))    #  502 920 EMU
FONT_SZ_100 = 1700       # sz attribute = hundredths of a point  (17 pt)

SRC  = '/mnt/c/Users/amits/Desktop/ultimate_presentation_v2.pptx'
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v3.pptx',
    os.path.join(os.path.dirname(os.path.abspath(__file__)), '..',
                 'outputs', 'ultimate_presentation_v3.pptx'),
]


def _set_xfrm(sp_el, l, t, w, h):
    """Set or create <a:xfrm> inside <p:spPr> with given EMU values."""
    spPr = sp_el.find(qn('p:spPr'))
    if spPr is None:
        spPr = etree.SubElement(sp_el, qn('p:spPr'))

    xfrm = spPr.find(qn('a:xfrm'))
    if xfrm is None:
        xfrm = etree.SubElement(spPr, qn('a:xfrm'))
        # insert at position 0 so it precedes prstGeom etc.
        spPr.insert(0, xfrm)

    off = xfrm.find(qn('a:off'))
    if off is None:
        off = etree.SubElement(xfrm, qn('a:off'))
    off.set('x', str(l))
    off.set('y', str(t))

    ext = xfrm.find(qn('a:ext'))
    if ext is None:
        ext = etree.SubElement(xfrm, qn('a:ext'))
    ext.set('cx', str(w))
    ext.set('cy', str(h))


def _set_font_size(sp_el, sz_100):
    """Set sz on every <a:rPr> run in the shape's txBody to sz_100 (100ths pt)."""
    txBody = sp_el.find(qn('p:txBody'))
    if txBody is None:
        return
    # also set body-level default via lstStyle / a:lvl1pPr / a:defRPr
    lst = txBody.find(qn('a:lstStyle'))
    if lst is not None:
        lvl = lst.find(qn('a:lvl1pPr'))
        if lvl is not None:
            defRPr = lvl.find(qn('a:defRPr'))
            if defRPr is not None:
                defRPr.set('sz', str(sz_100))

    for rPr in txBody.iter(qn('a:rPr')):
        rPr.set('sz', str(sz_100))
        rPr.set('b', '1')   # ensure bold (matching figure slide title style)


prs = Presentation(SRC)
n_changed = 0

for i, sl in enumerate(prs.slides):
    # Only process story slides
    if sl.background.fill.type == MSO_FILL.SOLID:
        continue

    for s in sl.shapes:
        el = s.element
        ph = el.find('.//' + qn('p:ph'))
        if ph is None:
            continue
        ph_type = ph.get('type', 'body')
        if ph_type != 'title':          # skip ctrTitle, body, subTitle, sldNum
            continue

        _set_xfrm(el, L, T, W, H)
        _set_font_size(el, FONT_SZ_100)
        n_changed += 1
        print(f'  Slide {i+1:02d}: repositioned title "{s.text[:50]}"')

print(f'\n{n_changed} title shapes updated.')

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved → {out}')
print('Done.')
