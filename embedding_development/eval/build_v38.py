#!/usr/bin/env python3
"""
build_v38.py  —  Fix S39, S50, S52.

Base: user's Desktop ultimate_presentation_v37.pptx
v37 (91) → v38 (91 slides, same count):
  S39: replace gpv_task_ensembles.png (heatmap only, no bar chart)
       + patch title to remove 'Task-Sensitive Ensembles Highlighted'
  S50: shorten title (was getting cut off)
  S52: replace head_angle_attribution_summary.png (uniform green/orange + box)
"""
import os, zipfile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.parts.presentation import PresentationPart


@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
mdir = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')

SRC  = '/mnt/c/Users/amits/Desktop/ultimate_presentation_v37.pptx'
OUTS = [
    os.path.join(root, 'outputs', 'ultimate_presentation_v38.pptx'),
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v38.pptx',
]

TARGET_DPI = 200


def verify_figure(path, label):
    if not os.path.exists(path):
        print(f'  ✗ MISSING: {label}'); return False
    with Image.open(path) as im:
        w, h = im.size
    print(f'  ✓ {label}: {w}×{h}px = {w/TARGET_DPI:.2f}"×{h/TARGET_DPI:.2f}"')
    return True


def verify_zip(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        dupes = [n for n in set(names) if names.count(n) > 1]
        slides = [n for n in names if 'slides/slide' in n
                  and n.endswith('.xml') and 'rels' not in n]
    ok = len(dupes) == 0
    print(f'  {"✓" if ok else "✗ DUPES:"+str(dupes)} zip: {len(slides)} slides, {len(dupes)} dupes')
    return ok


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def replace_picture(slide, new_img, left_in, top_in, w_in, h_in):
    for sh in list(slide.shapes):
        if sh.shape_type == 13:
            sh._element.getparent().remove(sh._element)
            break
    slide.shapes.add_picture(new_img,
        int(Inches(left_in)), int(Inches(top_in)),
        int(Inches(w_in)),    int(Inches(h_in)))


def patch_title(slide, old_fragment, new_text):
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            full = ''.join(r.text for r in para.runs)
            if old_fragment.lower() in full.lower():
                if para.runs:
                    para.runs[0].text = new_text
                    for r in list(para.runs)[1:]:
                        r.text = ''
                return True
    return False


# ── Pre-build verification ─────────────────────────────────────────────────────
print('PRE-BUILD VERIFICATION')
figs = [
    (os.path.join(mdir, 'gpv_task_ensembles.png'),          'GPV heatmap'),
    (os.path.join(mdir, 'head_angle_attribution_summary.png'), 'HA attribution'),
]
all_ok = all(verify_figure(p, l) for p, l in figs)
all_ok &= os.path.exists(SRC)
if not all_ok:
    raise SystemExit('Verification failed')
print('✓ All figures present\n')

# ── Build ──────────────────────────────────────────────────────────────────────
prs = Presentation(SRC)
print(f'Opened {os.path.basename(SRC)}: {len(prs.slides)} slides')

# ── S39: heatmap-only figure + title fix ──────────────────────────────────────
idx = find_idx(prs, 'Task-Sensitive Ensembles')
if idx is not None:
    replace_picture(prs.slides[idx],
                    os.path.join(mdir, 'gpv_task_ensembles.png'),
                    0.25, 0.62, 9.50, 4.20)
    patched = patch_title(prs.slides[idx],
                          'Task-Sensitive Ensembles Highlighted',
                          'GPV Heatmap: Feature Attribution by Ensemble')
    print(f'  S{idx+1}: GPV heatmap replaced, title patched={patched}')

# ── S50: shorten title ────────────────────────────────────────────────────────
idx = find_idx(prs, 'Cross-Ensemble Neural Prediction')
if idx is not None:
    patched = patch_title(prs.slides[idx],
                          'Embedding Consistency: Cross-Ensemble Neural Prediction',
                          'Cross-Ensemble Prediction')
    print(f'  S{idx+1}: title shortened, patched={patched}')

# ── S52: head angle attribution summary (green/orange + box) ──────────────────
idx = find_idx(prs, 'Head Angle Is the Top-Attributed')
if idx is not None:
    replace_picture(prs.slides[idx],
                    os.path.join(mdir, 'head_angle_attribution_summary.png'),
                    0.25, 0.60, 9.50, 3.50)
    print(f'  S{idx+1}: head angle attribution summary replaced')

print(f'\nFinal slide count: {len(prs.slides)}')

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved → {out}')

print('\nPOST-BUILD VERIFICATION')
for out in OUTS:
    if os.path.exists(out):
        verify_zip(out)
