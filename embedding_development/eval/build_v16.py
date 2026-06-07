#!/usr/bin/env python3
"""
build_v16.py  —  Apply remaining corrections to v15 (68 slides).

Changes:
  1. S26 (idx 25): replace r2_comparison_grand_mean.png  (internal title removed)
  2. S36 (idx 35): title → "Encoding Input Importance"
  3. S47 (idx 46): replace embedding_consistency_comparison.png (standard height, no violin)
  4. S48 (idx 47): merge cross-model + cross-ensemble attribution consistency → 2-panel slide
                   (replaces S48 figure, deletes S51)
  5. S49 (idx 48): replace embedding_linear_map.png (no n= labels, simpler 4-group layout)
                   + fix subtitle text (remove old decodability description)
  6. S50 (idx 49): replace cross_ensemble_prediction.png (violin only, no heatmap)
  7. S61 (idx 60): replace position_ablation.png (bottom margin fixed)

Net: 68 - 1 (delete S51) = 67 slides
"""
import os, zipfile, xml.etree.ElementTree as ET
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
import matplotlib; matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.gridspec as gridspec
import numpy as np

base  = os.path.dirname(os.path.abspath(__file__))
root  = os.path.join(base, '..')
mdir  = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')
cdir  = os.path.join(root, 'outputs', 'cebra_comparison')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v15.pptx')
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v16.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation_v16.pptx'),
]

SLIDE_W = Inches(10.0)
SLIDE_H = Inches(5.625)
DPI     = 200


def _img_wh(path):
    img = Image.open(path)
    return img.width / DPI, img.height / DPI


def _fit(fw, fh, max_w=Inches(9.5), max_h=Inches(4.8), top=Inches(0.75)):
    scale = min(max_w / Inches(fw), max_h / Inches(fh))
    w = int(Inches(fw) * scale); h = int(Inches(fh) * scale)
    l = int((SLIDE_W - w) // 2); t = int(top + (max_h - h) // 2)
    return l, t, w, h


def replace_figure(slide, img_path):
    pics = [s for s in slide.shapes if s.shape_type == 13]
    if not pics:
        print(f'  WARNING: no picture on slide')
        return
    target = max(pics, key=lambda s: s.width * s.height)
    l, t, w, h = int(target.left), int(target.top), int(target.width), int(target.height)
    target._element.getparent().remove(target._element)
    slide.shapes.add_picture(img_path, l, t, w, h)


def set_run(slide, match_fn, new_text):
    for sh in slide.shapes:
        try:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    if match_fn(run.text):
                        run.text = new_text
                        return True
        except Exception:
            pass
    return False


def remove_text_box(slide, match_fn):
    """Remove first text box whose full text satisfies match_fn."""
    for sh in list(slide.shapes):
        try:
            t = sh.text_frame.text.strip()
            if match_fn(t):
                sh._element.getparent().remove(sh._element)
                return True
        except Exception:
            pass
    return False


def delete_slide(prs, idx):
    """Remove slide at idx from sldIdLst AND its relationship from presentation.xml.rels."""
    from pptx.oxml.ns import qn
    xml_slides = prs.slides._sldIdLst
    children   = list(xml_slides)
    el         = children[idx]
    # Get the rId before removing from sldIdLst
    NS_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    rId  = el.get(f'{{{NS_R}}}id')
    xml_slides.remove(el)
    # Also remove from presentation.xml.rels so no orphaned relationship remains
    prs_part = prs.part
    if rId in prs_part._rels:
        del prs_part._rels[rId]


# ══════════════════════════════════════════════════════════════════════════════
# Build merged attribution consistency figure (S48)
# Left:  cross-model (MLP vs TC-Cont, MLP vs TC-Pred)
# Right: cross-ensemble (within-session, different ensembles)
# ══════════════════════════════════════════════════════════════════════════════
import sys
sys.path.insert(0, root)
from utils.figure_style import FONT, apply_style, add_footnote, savefig_manifest
from scipy.stats import spearmanr
import pickle

def _build_merged_attribution_fig():
    mlp_gpv = np.load(os.path.join(mdir, 'importance_global_pv_semantic_seed42.npy'),
                      allow_pickle=True)
    mlp_r2  = np.nanmean(np.load(os.path.join(mdir, 'all_r2.npy')), axis=0)
    cc_gpv  = np.nanmean([
        np.load(os.path.join(root, 'outputs', 'cebra_64d_eval', 'ensembles',
                             f'importance_global_pv_semantic_seed{s}.npy'), allow_pickle=True)
        for s in [42, 43, 44, 45, 46] if os.path.exists(
            os.path.join(root, 'outputs', 'cebra_64d_eval', 'ensembles',
                         f'importance_global_pv_semantic_seed{s}.npy'))
    ], axis=0)
    cp_gpv  = np.nanmean([
        np.load(os.path.join(root, 'outputs', 'cebra_pred_64d_eval', 'ensembles',
                             f'importance_global_pv_semantic_seed{s}.npy'), allow_pickle=True)
        for s in [42, 43, 44, 45, 46] if os.path.exists(
            os.path.join(root, 'outputs', 'cebra_pred_64d_eval', 'ensembles',
                         f'importance_global_pv_semantic_seed{s}.npy'))
    ], axis=0)
    cc_r2 = np.nanmean(np.load(os.path.join(root, 'outputs', 'cebra_64d_eval',
                                             'ensembles', 'all_r2.npy')), axis=0)
    cp_r2 = np.nanmean(np.load(os.path.join(root, 'outputs', 'cebra_pred_64d_eval',
                                             'ensembles', 'all_r2.npy')), axis=0)

    R2_THR = 0.1
    n_sess, n_ens = mlp_r2.shape

    def cross_rhos(gpv_a, r2_a, gpv_b, r2_b):
        rhos = []
        for s in range(n_sess):
            for e in range(n_ens):
                if r2_a[s, e] < R2_THR or r2_b[s, e] < R2_THR: continue
                va = gpv_a[s, e]; vb = gpv_b[s, e]
                ok = np.isfinite(va) & np.isfinite(vb)
                if ok.sum() < 4: continue
                r, _ = spearmanr(va[ok], vb[ok])
                rhos.append(r)
        return np.array(rhos)

    # Cross-ensemble: within session, different ensembles
    ce_rhos = []
    for s in range(n_sess):
        for e1 in range(n_ens):
            if mlp_r2[s, e1] < R2_THR: continue
            for e2 in range(e1 + 1, n_ens):
                if mlp_r2[s, e2] < R2_THR: continue
                va = mlp_gpv[s, e1]; vb = mlp_gpv[s, e2]
                ok = np.isfinite(va) & np.isfinite(vb)
                if ok.sum() < 4: continue
                r, _ = spearmanr(va[ok], vb[ok])
                ce_rhos.append(r)
    ce_rhos = np.array(ce_rhos)

    fig, (ax_l, ax_r) = plt.subplots(1, 2, figsize=(9.5, 4.2))
    apply_style(fig, ax_l); apply_style(fig, ax_r)
    rng = np.random.default_rng(0)

    # Left: cross-model
    pairs = [
        ('MLP\nvs CC', cross_rhos(mlp_gpv, mlp_r2, cc_gpv, cc_r2), '#4CAF50'),
        ('MLP\nvs CP', cross_rhos(mlp_gpv, mlp_r2, cp_gpv, cp_r2), '#FF9800'),
        ('CC\nvs CP',  cross_rhos(cc_gpv,  cc_r2,  cp_gpv, cp_r2), '#2196F3'),
    ]
    vp = ax_l.violinplot([p[1] for p in pairs], positions=[0,1,2],
                          showmedians=True, showextrema=True, widths=0.55)
    for pc, (_, _, c) in zip(vp['bodies'], pairs): pc.set_facecolor(c); pc.set_alpha(0.45)
    for part in ['cmedians','cmins','cmaxes','cbars']: vp[part].set_color('#333'); vp[part].set_linewidth(0.9)
    for i, (lbl, rhos, c) in enumerate(pairs):
        ax_l.scatter(i+rng.uniform(-0.08,0.08,len(rhos)), rhos, s=6, color=c, alpha=0.4, linewidths=0, zorder=3)
        ax_l.text(i, np.median(rhos)+0.02, f'{np.median(rhos):.2f}',
                  ha='center', va='bottom', fontsize=FONT.ANNOTATION-1, fontweight='bold')
    ax_l.set_xticks([0,1,2]); ax_l.set_xticklabels([p[0] for p in pairs], fontsize=FONT.TICK)
    ax_l.set_ylabel('Spearman ρ of GPV profiles', fontsize=FONT.LABEL)
    ax_l.set_ylim(-0.1, 1.15); ax_l.axhline(0.9, color='#555', ls=':', lw=0.9)
    ax_l.set_title('Cross-architecture', fontsize=FONT.LABEL)

    # Right: cross-ensemble
    vp2 = ax_r.violinplot([ce_rhos], positions=[0],
                           showmedians=True, showextrema=True, widths=0.55)
    for pc in vp2['bodies']: pc.set_facecolor('#4CAF50'); pc.set_alpha(0.45)
    for part in ['cmedians','cmins','cmaxes','cbars']: vp2[part].set_color('#333'); vp2[part].set_linewidth(0.9)
    ax_r.scatter(rng.uniform(-0.08,0.08,len(ce_rhos)), ce_rhos, s=6, color='#4CAF50', alpha=0.4, linewidths=0, zorder=3)
    ax_r.text(0, np.median(ce_rhos)+0.02, f'{np.median(ce_rhos):.2f}',
              ha='center', va='bottom', fontsize=FONT.ANNOTATION-1, fontweight='bold')
    ax_r.set_xticks([0]); ax_r.set_xticklabels(['MLP\n(diff. ensembles)'], fontsize=FONT.TICK)
    ax_r.set_ylabel('Spearman ρ of GPV profiles', fontsize=FONT.LABEL)
    ax_r.set_ylim(-0.1, 1.15); ax_r.axhline(0.9, color='#555', ls=':', lw=0.9)
    ax_r.set_title('Cross-ensemble (same session)', fontsize=FONT.LABEL)

    add_footnote(fig, f'GPV Spearman ρ; R²≥{R2_THR}; CC=TC-Cont CP=TC-Pred')
    savefig_manifest(fig, 'attribution_consistency_merged.png',
                     [mdir, cdir, '/mnt/c/Users/amits/Desktop/analysis_figures'])
    print('  Built attribution_consistency_merged.png')

print('Building merged attribution figure...')
try:
    _build_merged_attribution_fig()
except Exception as e:
    print(f'  WARNING: could not build merged fig: {e}')

# ══════════════════════════════════════════════════════════════════════════════
# Open and edit PPTX
# ══════════════════════════════════════════════════════════════════════════════
prs = Presentation(SRC)
print(f'\nOpened v15: {len(prs.slides)} slides')

# 1. S26 (idx 25) — new grand mean bars (no internal title)
img = os.path.join(cdir, 'r2_comparison_grand_mean.png')
if os.path.exists(img):
    replace_figure(prs.slides[25], img)
    print('  S26: replaced r2_comparison_grand_mean.png')

# 2. S36 (idx 35) — title change
ok = set_run(prs.slides[35],
             lambda t: 'What Does the Best Model' in t or 'Feature Attribution' in t,
             'Encoding Input Importance')
print(f'  S36 title: {"updated" if ok else "NOT FOUND"}')

# 3. S47 (idx 46) — new heatmap (standard height)
img = os.path.join(cdir, 'embedding_consistency_comparison.png')
if os.path.exists(img):
    replace_figure(prs.slides[46], img)
    print('  S47: replaced embedding_consistency_comparison.png')

# 4. S48 (idx 47) — merged attribution consistency figure
merged_img = os.path.join(mdir, 'attribution_consistency_merged.png')
if os.path.exists(merged_img):
    replace_figure(prs.slides[47], merged_img)
    ok = set_run(prs.slides[47],
                 lambda t: 'Attribution Consistency' in t or 'Cross-Model' in t,
                 'Attribution Profile Consistency: Cross-Architecture and Cross-Ensemble')
    print(f'  S48: replaced with merged figure, title {"updated" if ok else "not found"}')

# 5. S49 (idx 48) — new embedding_linear_map + fix subtitle
img = os.path.join(mdir, 'embedding_linear_map.png')
if os.path.exists(img):
    replace_figure(prs.slides[48], img)
    print('  S49: replaced embedding_linear_map.png')
# Remove old decodability subtitle
removed = remove_text_box(prs.slides[48],
    lambda t: 'decodability' in t.lower() or 'Spearman' in t or 'ridge probe' in t.lower())
print(f'  S49 old subtitle: {"removed" if removed else "not found"}')
# Update title
ok = set_run(prs.slides[48],
             lambda t: 'Representation Consistency' in t or 'Geometric Consistency' in t,
             'Embedding Geometric Consistency: Linear Map R² Across Variation Axes')
print(f'  S49 title: {"updated" if ok else "not found (may already be correct)"}')

# 6. S50 (idx 49) — violin-only cross_ensemble_prediction
img = os.path.join(mdir, 'cross_ensemble_prediction.png')
if os.path.exists(img):
    replace_figure(prs.slides[49], img)
    # Also fix subtitle
    set_run(prs.slides[49],
            lambda t: 'Self vs cross-target R²' in t or 'source×target' in t,
            'Self-target vs cross-target prediction R² by architecture (source models R²≥0.1)')
    print('  S50: replaced cross_ensemble_prediction.png')

# 7. S61 (idx 60) — new position_ablation (bottom margin)
pos_img_path = os.path.join(mdir, 'position_ablation.png')
if not os.path.exists(pos_img_path):
    pos_img_path = os.path.join(root, 'outputs', 'mlps', 'head_angle_analysis', 'position_ablation.png')
if os.path.exists(pos_img_path):
    replace_figure(prs.slides[60], pos_img_path)
    print('  S61: replaced position_ablation.png')
else:
    print(f'  S61: SKIP (position_ablation.png not found at expected path)')

# 8. Delete S51 (now redundant — merged into S48)
# After S50 fix, S51 is at idx 50
# Verify it's cross-ensemble attribution
s51_text = ''
for sh in prs.slides[50].shapes:
    try: s51_text += sh.text_frame.text
    except: pass
if 'Cross-Ensemble Attribution' in s51_text or 'cross_ensemble_consistency' in s51_text:
    delete_slide(prs, 50)
    print('  S51: deleted (merged into S48)')
else:
    print(f'  S51 skip: unexpected content: {s51_text[:60]}')

# ── Save ──────────────────────────────────────────────────────────────────────
print(f'\nFinal slide count: {len(prs.slides)}')
for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved  → {out}')
