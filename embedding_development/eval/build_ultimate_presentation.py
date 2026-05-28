#!/usr/bin/env python3
"""
build_ultimate_presentation.py

Combines story slides from Neural_Task_Representation_cleanup_v3.pptx with
all new figures from the manifest to produce the ultimate presentation.

Strategy: open cleanup_v3 as the working prs, append all new figure slides
(they inherit the dark master), then reorder _sldIdLst to produce the desired
sequence. Slides not in the desired order are implicitly dropped.
Old-slide indices below are 0-based.
"""

import os, sys, copy, shutil
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from PIL import Image

# ── Paths ──────────────────────────────────────────────────────────────────────
SRC_PPTX = '/mnt/c/Users/amits/Downloads/Neural_Task_Representation_cleanup_v3.pptx'

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')

FIGURE_DIRS = [
    os.path.join(root, 'outputs', 'cebra_comparison'),
    os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed'),
    os.path.join(root, 'outputs', 'mlps', 'ml_vs_naive'),
    os.path.join(root, 'outputs', 'cebra_eval', 'ensembles'),
    os.path.join(root, 'outputs', 'temporal_advantage'),
]
DPI = 200

OUT_PATHS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation.pptx'),
]

# ── Figure index ───────────────────────────────────────────────────────────────
figs = {}
for d in FIGURE_DIRS:
    if os.path.isdir(d):
        for fn in os.listdir(d):
            if fn.endswith('.png'):
                figs[fn] = os.path.join(d, fn)
print(f"Found {len(figs)} figure files")

# ── Slide plan ─────────────────────────────────────────────────────────────────
# ('old', 0-based-idx)  — copy from cleanup_v3
# ('fig', filename, title_text) — new figure slide
SLIDE_PLAN = [
    # ─ INTRO ──────────────────────────────────────────────────────────────────
    ('old', 0),   # Title
    ('old', 1),   # Example trial (has MP4 video)
    ('old', 2),   # Paradigm & Setup
    ('old', 3),   # Data Distribution
    # ─ DATASET ────────────────────────────────────────────────────────────────
    ('old', 5),   # Input Data Features: P(s, a)  [slide 6]
    ('old', 6),   # Input Data Features: Continued [slide 7]
    ('fig', 'behavioral_trace.png',
     'Behavioral Features: All 11 Signals for One Representative Trial'),
    # ─ METHODS ────────────────────────────────────────────────────────────────
    ('old', 4),   # Initial idea: SIPEC → Ensemble [slide 5]
    ('old', 7),   # Recover bin order [slide 8]
    ('old', 8),   # Approach Overview [slide 9]
    ('old', 9),   # Dataset Details [slide 10]
    ('old', 10),  # Predictive Encoding Metrics [slide 11]
    # ─ MODELS & R² ────────────────────────────────────────────────────────────
    ('old', 11),  # Linear Model: affine? [slide 12]
    ('fig', 'r2_bar_linear.png',
     'Linear Baseline: Within-Session Ensemble R² (29 sessions × 23 ensembles)'),
    ('old', 14),  # MLP Model: non-linear [slide 15]
    ('old', 15),  # MLP Capability Demo (GIFs) [slide 16]
    ('fig', 'r2_bar_mlp.png',
     'MLP Baseline: Within-Session Ensemble R²'),
    ('fig', 'r2_grand_mean_bars.png',
     'Grand Mean R² Across All Four Model Architectures'),
    ('fig', 'two_thresholds_scatter.png',
     'Valid (R²≥0.01 / R²≥0.05) Pair Counts per Session: MLP vs Linear'),
    ('fig', 'r2_consistency_bar.png',
     'MLP Cross-Seed Consistency (Pearson r on Held-out Trials)'),
    ('fig', 'embedding_consistency_comparison.png',
     'Embedding Consistency Across Models: MLP, TempConv-Cont, TempConv-Pred'),
    # ─ ATTRIBUTION ────────────────────────────────────────────────────────────
    ('old', 18),  # Permutation Variance explanation [slide 19]
    ('fig', 'gpv_group_ensemble_heatmap.png',
     'Global Permutation Variance by Feature Group × Ensemble (sorted by R²)'),
    ('old', 26),  # Conditional Permutation Variance [slide 27]
    ('fig', 'global_vs_cond_pv_scatter.png',
     'Global PV vs Conditional PV — Points Below Diagonal Are Collinearity-Inflated'),
    ('old', 29),  # Integrated Gradients explanation [slide 30]
    ('fig', 'ig_per_ensemble_heatmap.png',
     'Mean |IG| Attribution by Feature Group × Ensemble'),
    ('fig', 'group_attribution_comparison.png',
     'Group-Level Attribution Profile: MLP vs TempConv-Cont (GPV and IG)'),
    ('fig', 'attribution_agreement_scatter.png',
     'Attribution Agreement: MLP vs TempConv-Cont (GPV ρ≈0.95, IG ρ≈0.94)'),
    # ─ CASE STUDIES ───────────────────────────────────────────────────────────
    ('old', 32),  # Case Study E07 & E23 setup slide [slide 33]
    ('fig', 'case_studies_e07_e23.png',
     'Case Studies: E07 (Speed-Dominated) vs E23 (Head-Angle Dominated)'),
    # ─ FREQUENCY CAPABILITY ───────────────────────────────────────────────────
    ('old', 38),  # MLP Fails on High Frequency [slide 39]
    ('fig', 'trend_noise_comparison.png',
     'MLP Fails on Fast Neural Fluctuations; TempConv-Pred Captures Them'),
    # ─ TEMPCONV / CEBRA ───────────────────────────────────────────────────────
    ('old', 39),  # Temporal Context and Contrastive Learning [slide 40]
    ('old', 40),  # Can Contrastive Embeddings Reveal Behavioral Structure? [slide 41]
    # ─ HEAD ANGLE CASE STUDIES ────────────────────────────────────────────────
    ('old', 35),  # Head Angle Analysis [slide 36]
    ('old', 36),  # Head Angle PCA Animation — GIFs! [slide 37]
    ('fig', 'head_angle_scatter_2x2.png',
     'Top Pairs: Head Angle × Head Angular Velocity (Colored by z-Scored Activation)'),
    ('fig', 'head_angle_stability.png',
     'E18 Tuning Curve Stability — Preferred Angle Consistent Across Sessions'),
    ('fig', 'head_angle_tuning_6panel.png',
     'E18 Tuning Shape Variety Across Six Representative Sessions'),
    ('fig', 'position_tuning.png',
     'Position Tuning Curves for Top 4 Pairs (MLP Cannot Decode Position via Attribution)'),
    # ─ ML vs NAIVE ────────────────────────────────────────────────────────────
    ('fig', 'ml_vs_naive_scatter.png',
     'Part 1 — Validation: ML Attribution Tracks Naive Data-Effect Size'),
    ('fig', 'nonlinearity_advantage.png',
     'Part 2 — Nonlinearity: ML Captures Tuning That Spearman ρ Misses'),
    ('fig', 'mlp_vs_linear_r2.png',
     'Part 3 — ML vs GLM: MLP Outperforms Linear for Attribution-Flagged Pairs'),
    ('fig', 'ablation_proof.png',
     'Part 4 — Ablation Proof: Single-Feature MLP Predicts Where GLM Sees Nothing'),
    # ─ SAMPLE PREDICTED TRAJECTORIES ─────────────────────────────────────────
    ('old', 49),  # Example Plots 1 [slide 50]
    ('old', 50),  # Example Plots 2 [slide 51]
    ('old', 51),  # Example Plots 3 [slide 52]
]

# ── Style constants ─────────────────────────────────────────────────────────────
SLIDE_W  = Inches(10.0)
SLIDE_H  = Inches(5.625)
TITLE_H  = Inches(0.55)
DARK_BG  = RGBColor(0x1A, 0x23, 0x3A)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

def _set_dark_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = DARK_BG

def _add_textbox(slide, left, top, width, height, text, font_size,
                 bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p  = tf.paragraphs[0]
    p.alignment = align
    r  = p.add_run()
    r.text = text
    r.font.size      = font_size
    r.font.bold      = bold
    r.font.color.rgb = color
    return tb

def add_figure_slide(prs, filename, title_text):
    """Append a new dark-background figure slide to prs."""
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    _set_dark_bg(slide)

    _add_textbox(slide,
        Inches(0.15), Inches(0.05),
        SLIDE_W - Inches(0.3), TITLE_H,
        title_text, Pt(17), bold=True, color=WHITE)

    if filename not in figs:
        print(f"    [MISSING] {filename}")
        return len(prs.slides) - 1

    img_path  = figs[filename]
    img       = Image.open(img_path)
    fig_w     = img.width  / DPI
    fig_h     = img.height / DPI

    content_top = TITLE_H + Inches(0.05)
    content_h   = (SLIDE_H - content_top - Inches(0.08)) / 914400
    content_w   = (SLIDE_W - Inches(0.10)) / 914400

    scale    = min(content_w / fig_w, content_h / fig_h)
    placed_w = Inches(fig_w * scale)
    placed_h = Inches(fig_h * scale)
    left     = (SLIDE_W - placed_w) // 2
    top      = content_top + Emu(int((content_h - fig_h * scale) / 2 * 914400))

    slide.shapes.add_picture(img_path, left, top, placed_w, placed_h)

    # Filename caption bottom-right
    cap = slide.shapes.add_textbox(
        SLIDE_W - Inches(4.0), SLIDE_H - Inches(0.22),
        Inches(3.9), Inches(0.20))
    p2 = cap.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2 = p2.add_run()
    r2.text = filename
    r2.font.size      = Pt(9)
    r2.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    return len(prs.slides) - 1   # 0-based index of the newly added slide


# ── Build ───────────────────────────────────────────────────────────────────────
print(f"\nOpening source: {SRC_PPTX}")
prs = Presentation(SRC_PPTX)
n_original = len(prs.slides)
print(f"  {n_original} original slides")

# Step 1: Append all new figure slides (preserving their order for fig_idx_map)
fig_idx_map = {}   # filename -> 0-based slide index in prs
for entry in SLIDE_PLAN:
    if entry[0] == 'fig':
        fname = entry[1]
        if fname not in fig_idx_map:
            title = entry[2]
            idx   = add_figure_slide(prs, fname, title)
            fig_idx_map[fname] = idx
            print(f"  + Added figure slide [{idx}]: {fname}")

print(f"\n  {len(prs.slides)} total slides after adding figures")

# Step 2: Build desired 0-based index sequence
desired_order = []
for entry in SLIDE_PLAN:
    if entry[0] == 'old':
        desired_order.append(entry[1])
    else:
        desired_order.append(fig_idx_map[entry[1]])

print(f"\n  Desired order ({len(desired_order)} slides): {desired_order[:10]}...")

# Step 3: Reorder _sldIdLst
xml_slides  = prs.slides._sldIdLst
all_sl_els  = list(xml_slides)           # snapshot of all slide XML elements

for el in list(xml_slides):             # clear current list
    xml_slides.remove(el)

for idx in desired_order:              # add back in desired order
    xml_slides.append(all_sl_els[idx])

print(f"  Final slide count: {len(prs.slides)}")

# Step 4: Save
for out_path in OUT_PATHS:
    os.makedirs(os.path.dirname(out_path) or '.', exist_ok=True)
    prs.save(out_path)
    print(f"\nSaved ({len(prs.slides)} slides) → {out_path}")

print("\nDone.")
