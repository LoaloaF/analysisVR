#!/usr/bin/env python3
"""
build_v8.py  —  Apply figure updates to v7 → v8.

Changes from v7 (52 slides):
  - S22: replace r2_grand_mean_bars.png  (no more SD bars)
  - S24: replace two_thresholds_scatter.png  (single panel, R²≥0.05 only)
  - S30: replace gpv_group_ensemble_heatmap.png  (round colorbar ticks)
  - S34: replace ig_per_ensemble_heatmap.png  (round colorbar ticks)
  - S37: replace joint_effect_scatter.png  (η² instead of ρ)
  - S45: replace head_angle_scatter_2x2.png  (z-scored axis labels)
  - S49: replace position_collinearity.png  (fixed position label, dynamic features)

  + Insert new slide after S21: r2_bar_tempconv_comparison.png
    → shifts S22→S23, S24→S25, S30→S31, S34→S35, S37→S38, S45→S46, S49→S50

  + Add hidden slides at end for methods reference.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image

SRC  = '/mnt/c/Users/amits/Desktop/ultimate_presentation_v7.pptx'
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v8.pptx',
    os.path.join(os.path.dirname(os.path.abspath(__file__)), '..',
                 'outputs', 'ultimate_presentation_v8.pptx'),
]

base    = os.path.dirname(os.path.abspath(__file__))
root    = os.path.join(base, '..')
mdir    = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')
cdir    = os.path.join(root, 'outputs', 'cebra_comparison')

SLIDE_W = Inches(10.0)
SLIDE_H = Inches(5.625)
TITLE_H = Inches(0.55)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
DARK    = RGBColor(0x1A, 0x23, 0x3A)
GREY    = RGBColor(0x55, 0x55, 0x55)
DPI     = 200


def _img_size(img_path):
    img = Image.open(img_path)
    return img.width / DPI, img.height / DPI   # inches


def _fit_in_content_area(fw, fh):
    """Return (left, top, pw, ph) Emu values fitting image in slide content area."""
    ct = TITLE_H + Inches(0.05)
    ch = (SLIDE_H - ct - Inches(0.08)) / 914400   # inches
    cw = (SLIDE_W - Inches(0.10)) / 914400
    sc = min(cw / fw, ch / fh)
    pw = Inches(fw * sc)
    ph = Inches(fh * sc)
    left = (SLIDE_W - pw) // 2
    top  = ct + Emu(int((ch - fh * sc) / 2 * 914400))
    return left, top, pw, ph


def replace_figure(slide, new_img_path):
    """Remove the existing picture shape(s) and add the new image, same content area."""
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            shape._element.getparent().remove(shape._element)
    fw, fh = _img_size(new_img_path)
    left, top, pw, ph = _fit_in_content_area(fw, fh)
    slide.shapes.add_picture(new_img_path, left, top, pw, ph)


def add_figure_slide(prs, img_path, title_text):
    layout = prs.slide_layouts[6]
    slide  = prs.slides.add_slide(layout)
    fill   = slide.background.fill
    fill.solid(); fill.fore_color.rgb = WHITE

    tb = slide.shapes.add_textbox(Inches(0.15), Inches(0.05),
                                   SLIDE_W - Inches(0.3), TITLE_H)
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title_text
    r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = DARK

    fw, fh = _img_size(img_path)
    left, top, pw, ph = _fit_in_content_area(fw, fh)
    slide.shapes.add_picture(img_path, left, top, pw, ph)

    cap = slide.shapes.add_textbox(SLIDE_W - Inches(4.0), SLIDE_H - Inches(0.22),
                                    Inches(3.9), Inches(0.20))
    p2  = cap.text_frame.paragraphs[0]
    p2.alignment = PP_ALIGN.RIGHT
    r2  = p2.add_run()
    r2.text = os.path.basename(img_path)
    r2.font.size = Pt(9); r2.font.color.rgb = GREY
    return len(prs.slides) - 1


def move_slide(prs, from_idx, to_idx):
    xml_slides = prs.slides._sldIdLst
    children   = list(xml_slides)
    el = children.pop(from_idx)
    children.insert(to_idx, el)
    for c in list(xml_slides):
        xml_slides.remove(c)
    for c in children:
        xml_slides.append(c)


def add_hidden_text_slide(prs, title_text, body_lines):
    """Add a hidden reference slide with bullet-point text."""
    layout = prs.slide_layouts[6]
    slide  = prs.slides.add_slide(layout)
    fill   = slide.background.fill
    fill.solid(); fill.fore_color.rgb = WHITE

    # Title
    tb = slide.shapes.add_textbox(Inches(0.15), Inches(0.05),
                                   SLIDE_W - Inches(0.3), TITLE_H)
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title_text
    r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = DARK

    # Body
    tb2 = slide.shapes.add_textbox(Inches(0.4), Inches(0.65),
                                    Inches(9.2), Inches(4.8))
    tf  = tb2.text_frame
    tf.word_wrap = True
    for i, line in enumerate(body_lines):
        if i == 0:
            para = tf.paragraphs[0]
        else:
            para = tf.add_paragraph()
        run = para.add_run()
        run.text = line
        run.font.size = Pt(12 if not line.startswith('  ') else 11)
        run.font.color.rgb = DARK
        run.font.bold = line.startswith('▶')

    # Mark hidden (PowerPoint respects show="0" on the slide element)
    slide._element.set('show', '0')
    return slide


# ── Open v7 ───────────────────────────────────────────────────────────────────
prs = Presentation(SRC)
print(f'Opened v7: {len(prs.slides)} slides')

# ── Step 1: Replace figures in existing slides (before any insertions) ─────────
# v7 0-based indices:
REPLACEMENTS = {
    21: (os.path.join(cdir, 'r2_grand_mean_bars.png'),       'r2_grand_mean_bars'),
    23: (os.path.join(cdir, 'two_thresholds_scatter.png'),   'two_thresholds_scatter'),
    29: (os.path.join(mdir, 'gpv_group_ensemble_heatmap.png'),'gpv_group_ensemble_heatmap'),
    33: (os.path.join(mdir, 'ig_per_ensemble_heatmap.png'),  'ig_per_ensemble_heatmap'),
    36: (os.path.join(mdir, 'joint_effect_scatter.png'),     'joint_effect_scatter'),
    44: (os.path.join(mdir, 'head_angle_scatter_2x2.png'),   'head_angle_scatter_2x2'),
    48: (os.path.join(mdir, 'position_collinearity.png'),    'position_collinearity'),
}
for idx_0, (img_path, name) in REPLACEMENTS.items():
    replace_figure(prs.slides[idx_0], img_path)
    print(f'  Replaced S{idx_0+1:02d} with {name}')

# ── Step 2: Insert new TempConv bar slide after S21 (idx 21) ──────────────────
tc_img = os.path.join(cdir, 'r2_bar_tempconv_comparison.png')
new_idx = add_figure_slide(prs, tc_img,
                            'TempConv-Cont and TempConv-Pred R² Per Ensemble (Sorted by MLP R²)')
move_slide(prs, new_idx, 21)   # insert BEFORE current S22 → new slot 21 (0-based)
print(f'  Inserted TempConv comparison bar at position 22 (1-based)')

# ── Step 3: Update S37 slide title to reflect η² ──────────────────────────────
# After insertion above, old S37 (idx 36) → idx 37 (0-based)
s37_slide = prs.slides[37]
for shape in s37_slide.shapes:
    if shape.has_text_frame:
        tf = shape.text_frame
        for para in tf.paragraphs:
            for run in para.runs:
                if 'No Single Feature Predicts' in run.text or 'ML Captures Joint' in run.text:
                    run.text = 'ML Captures Joint Effects: No Single Feature Explains ≥5% Variance (η²)'
                    break

print('  Updated S38 title to mention η²')

# ── Step 4: Add hidden reference slides ───────────────────────────────────────
HIDDEN_SLIDES = [
    (
        '[Hidden] Methods Reference: Attribution Metrics',
        [
            '▶ GPV — Global Permutation Variance (ΔR²)',
            '  Measures how much R² drops when a feature group is randomly permuted.',
            '  Captures the total importance including shared variance with other features.',
            '',
            '▶ CPV — Conditional Permutation Variance (ΔR²)',
            '  Permutes a feature while conditioning on all others.',
            '  Removes collinearity inflation; sensitive to features with unique information.',
            '',
            '▶ IG — Integrated Gradients (Mean |IG|)',
            '  Gradient-based attribution from the input along a path from baseline.',
            '  Signed: positive IG = feature increases predicted activity.',
            '',
            '▶ Consistency: Spearman ρ between feature attribution profiles',
            '  Cross-seed ρ ≈ 0.95 (5 seeds), cross-model ρ ≈ 0.95 (MLP vs TempConv).',
        ]
    ),
    (
        '[Hidden] Methods Reference: Effect Size Metrics',
        [
            '▶ R² (coefficient of determination)',
            '  Fraction of variance in neural activity explained by the model.',
            '  R² = 1 − SS_res / SS_tot.  Computed on held-out trials.',
            '',
            '▶ η² (eta squared, binned ANOVA)',
            '  Fraction of variance explained by a single feature via a step-function',
            '  predictor (10 quantile bins).  Captures non-linear univariate associations.',
            '  η² = SS_between / SS_total.',
            '  We use max η² across all features as a conservative univariate baseline.',
            '',
            '▶ Spearman ρ (rank correlation)',
            '  Monotonic association between two variables, robust to outliers.',
            '  ρ = 1 − 6Σd²/(n(n²−1)). Used for attribution-profile similarity.',
            '',
            '▶ 1-CV (embedding consistency)',
            '  1 minus the coefficient of variation of pairwise Pearson r values.',
            '  High 1-CV → consistent embedding geometry across seeds/sessions.',
        ]
    ),
    (
        '[Hidden] Methods Reference: Models',
        [
            '▶ Linear Encoding Model',
            '  Ridge regression from z-scored behavioral features to z-scored ensemble activity.',
            '  Baseline: captures only linear/affine feature combinations.',
            '',
            '▶ MLP (Multi-Layer Perceptron)',
            '  2-layer fully connected network, ReLU activations, dropout.',
            '  Trained per (session, seed); R² evaluated on held-out trials.',
            '',
            '▶ TempConv-Cont (CEBRA Contrastive)',
            '  Temporal convolutional encoder trained with contrastive objective.',
            '  Uses behavioral labels as positive-pair signal.',
            '',
            '▶ TempConv-Pred (CEBRA Predictive)',
            '  Same architecture but trained with a predictive (non-contrastive) objective.',
            '  ΔR² vs MLP: mean = −0.012 (architecture-independent representations).',
            '',
            '▶ Ensemble Activity',
            '  Linear combination of single-unit spikes using SIPEC ensemble weights.',
            '  Target: z-scored ensemble activation across trials.',
        ]
    ),
]

for title, lines in HIDDEN_SLIDES:
    add_hidden_text_slide(prs, title, lines)
    print(f'  Added hidden slide: {title[:50]}')

# ── Save ──────────────────────────────────────────────────────────────────────
print(f'\nFinal slide count: {len(prs.slides)}')
for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved → {out}')
print('Done.')
