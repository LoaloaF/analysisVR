#!/usr/bin/env python3
"""
build_v25.py  —  5 fixes from user feedback, v24 (84) → v25 (84 slides, same count).

1. S50: fix subtitle bold=None → bold=False to match S36/S46 format exactly
2. S51: replace head_angle_attribution_summary.png (annotation moved out of bars)
3. S59: replace honest_attribution_example.png (hspace=0.90, no subplot overlap)
4. S61: shorten title to "Case Studies" (was too long)
5. S62: replace e07_e23_covariate_validation.png (new story: detection + attribution wins)
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
mdir = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')
adir = os.path.join(root, 'outputs', 'ablation_vs_attribution')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v24.pptx')
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v25.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation_v25.pptx'),
]

C_SUB = RGBColor(0x33, 0x33, 0x33)


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def replace_picture(slide, new_img, left_in, top_in, w_in, h_in):
    for sh in list(slide.shapes):
        if sh.shape_type == 13:
            sh._element.getparent().remove(sh._element)
            break
    slide.shapes.add_picture(new_img,
        int(Inches(left_in)), int(Inches(top_in)),
        int(Inches(w_in)),    int(Inches(h_in)))


def patch_run(slide, old, new):
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    return True
    return False


prs = Presentation(SRC)
print(f'Opened v24: {len(prs.slides)} slides')


# ── 1. S50 subtitle bold fix ──────────────────────────────────────────────────
idx = find_idx(prs, 'Head Angle Analysis')
if idx is not None:
    sl = prs.slides[idx]
    for sh in sl.shapes:
        if sh.has_text_frame and 'Dominant feature' in sh.text_frame.text:
            for para in sh.text_frame.paragraphs:
                for run in para.runs:
                    run.font.bold = False          # explicit, matches S36
                    run.font.color.rgb = C_SUB     # ensure colour is set too
            print(f'  S{idx+1}: fixed subtitle bold=False')


# ── 2. S51 head_angle_attribution_summary ─────────────────────────────────────
idx = find_idx(prs, 'Head Angle Is the Top-Attributed')
if idx is not None:
    replace_picture(prs.slides[idx],
                    os.path.join(mdir, 'head_angle_attribution_summary.png'),
                    0.25, 0.60, 9.50, 3.50)
    print(f'  S{idx+1}: replaced head_angle_attribution_summary.png')


# ── 3. S59 honest_attribution_example (hspace fixed) ─────────────────────────
idx = find_idx(prs, 'Position Appears Primary by Correlation')
if idx is not None:
    replace_picture(prs.slides[idx],
                    os.path.join(mdir, 'honest_attribution_example.png'),
                    0.25, 0.62, 9.50, 4.80)
    print(f'  S{idx+1}: replaced honest_attribution_example.png')


# ── 4. S61 title shortened ────────────────────────────────────────────────────
idx = find_idx(prs, 'E07: Speed-Dominated')
if idx is not None:
    sl = prs.slides[idx]
    # Title is split across runs: 'E07: Speed-Dominated' | '  ·  E23: Head Angle-Dominated'
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            runs = para.runs
            # collect all runs that form the title
            full = ''.join(r.text for r in runs)
            if 'E07' in full and 'E23' in full:
                # clear all runs except first, set first to short title
                if runs:
                    runs[0].text = 'Case Studies'
                    for r in list(runs)[1:]:
                        r.text = ''
                print(f'  S{idx+1}: title shortened to "Case Studies"')
                break


# ── 5. S62 new covariate validation figure ────────────────────────────────────
idx = find_idx(prs, 'Speed (E07) and Head Angle')
if idx is None:
    idx = find_idx(prs, 'Case Studies: E07')
if idx is not None:
    replace_picture(prs.slides[idx],
                    os.path.join(adir, 'e07_e23_covariate_validation.png'),
                    0.25, 0.62, 9.50, 4.50)
    patch_run(prs.slides[idx],
              'Speed (E07) and Head Angle (E23) Outpredict Cue and Choice Labels',
              'When Cue/Choice Discriminates, Our Model Detects It — '
              'And Our Primary Variable Explains More')
    print(f'  S{idx+1}: replaced e07_e23_covariate_validation.png + updated title')


# ── Report ────────────────────────────────────────────────────────────────────
print(f'\nFinal slide count: {len(prs.slides)}')
for frag, note in [
    ('Head Angle Analysis',               'S50 section divider'),
    ('Head Angle Is the Top-Attributed',  'S51 summary bar chart'),
    ('Position Appears Primary',          'S59 honest attribution'),
    ('Case Studies',                      'S61 section header'),
    ('When Cue/Choice Discriminates',     'S62 validation figure'),
]:
    i = find_idx(prs, frag)
    print(f'  {note:<35} → {"S"+str(i+1) if i is not None else "NOT FOUND"}')

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'\nSaved → {out}')
