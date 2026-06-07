#!/usr/bin/env python3
"""
build_v47.py — Refresh regenerated figures in the deck.

Base: outputs/ultimate_presentation_v46.pptx  (94 slides)
v46 (94) → v47 (94 slides):
  No slide count change — replace figure images on 11 slides where
  labels, titles, and footnotes were cleaned up (em dashes removed,
  overlapping panel labels moved inside axes, long titles shortened).

  S41  global_vs_cond_pv_scatter.png
  S43  ig_per_ensemble_heatmap.png
  S49  attribution_consistency_merged.png
  S50  embedding_linear_map.png
  S51  cross_ensemble_prediction.png
  S55  head_angle_tuning_variety.png
  S65  e07_e23_bucket_signal.png
  S66  e07_e23_bucket_gpv.png
  S67  e07_e23_bucket_covariation.png
  S70  e07_cue_zone_joint_gpv.png
  S71  e07_cue_zone_competition.png
"""
import os, zipfile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.parts.presentation import PresentationPart

@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v46.pptx')
OUTS = [
    os.path.join(root, 'outputs', 'ultimate_presentation_v47.pptx'),
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v47.pptx',
]

TARGET_DPI = 200
mdir = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')
adir = os.path.join(root, 'outputs', 'ablation_vs_attribution')

# (title_fragment, img_path, slide_label)
UPDATES = [
    ('Global PV vs Conditional PV',     os.path.join(mdir, 'global_vs_cond_pv_scatter.png'),      'S41'),
    ('Mean |IG| Attribution',           os.path.join(mdir, 'ig_per_ensemble_heatmap.png'),          'S43'),
    ('Attribution Profile Consistency', os.path.join(mdir, 'attribution_consistency_merged.png'),   'S49'),
    ('Embedding Geometric Consistency', os.path.join(mdir, 'embedding_linear_map.png'),             'S50'),
    ('Cross-Ensemble Prediction',       os.path.join(mdir, 'cross_ensemble_prediction.png'),        'S51'),
    ('Head angle tuning',               os.path.join(mdir, 'head_angle_tuning_variety.png'),        'S55'),
    ('Neural signal: E07',              os.path.join(adir, 'e07_e23_bucket_signal.png'),            'S65'),
    ('Bucket 1: model attributes',      os.path.join(adir, 'e07_e23_bucket_gpv.png'),              'S66'),
    ('Bucket 2: a co-varying feature',  os.path.join(adir, 'e07_e23_bucket_covariation.png'),      'S67'),
    ('Joint GPV in the cue zone',       os.path.join(adir, 'e07_cue_zone_joint_gpv.png'),          'S70'),
    ('Joint GPV still fails',           os.path.join(adir, 'e07_cue_zone_competition.png'),         'S71'),
]


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def replace_figure(slide, new_img_path, prs, label):
    """Replace largest picture on slide with new image, recentered horizontally."""
    pic_shape = None
    max_area  = 0
    for shape in slide.shapes:
        if shape.shape_type == 13:
            area = shape.width * shape.height
            if area > max_area:
                max_area  = area
                pic_shape = shape

    if pic_shape is None:
        print(f'  WARNING: no picture found on {label}')
        return False

    old_top = pic_shape.top

    with Image.open(new_img_path) as im:
        iw, ih = im.size
    img_w_emu = int(iw * 914400 / TARGET_DPI)
    img_h_emu = int(ih * 914400 / TARGET_DPI)
    left_emu  = (prs.slide_width - img_w_emu) // 2

    pic_shape._element.getparent().remove(pic_shape._element)
    slide.shapes.add_picture(new_img_path, left_emu, old_top, img_w_emu, img_h_emu)
    print(f'      replaced  {iw}×{ih}px  ({iw/TARGET_DPI:.1f}"×{ih/TARGET_DPI:.1f}")')
    return True


def verify_zip(path):
    with zipfile.ZipFile(path) as z:
        names  = z.namelist()
        dupes  = [n for n in set(names) if names.count(n) > 1]
        slides = [n for n in names if 'slides/slide' in n
                  and n.endswith('.xml') and 'rels' not in n]
    ok = len(dupes) == 0
    print(f'  {"ok" if ok else "DUPES: "+str(dupes)}  {len(slides)} slides  {len(dupes)} dupes')
    return ok


# ── Pre-build checks ─────────────────────────────────────────────────────────
print('PRE-BUILD CHECKS')
all_ok = True
for _, img_path, label in UPDATES:
    if os.path.exists(img_path):
        with Image.open(img_path) as im:
            print(f'  ok {label}: {os.path.basename(img_path)}  {im.size[0]}x{im.size[1]}px')
    else:
        print(f'  MISSING {label}: {img_path}')
        all_ok = False
if not all_ok:
    raise SystemExit('Missing figures — aborting')
if not os.path.exists(SRC):
    raise SystemExit(f'Source not found: {SRC}')
print(f'  ok source: {os.path.basename(SRC)}\n')

# ── Build ─────────────────────────────────────────────────────────────────────
prs = Presentation(SRC)
print(f'Opened {os.path.basename(SRC)}: {len(prs.slides)} slides\n')

print('REPLACING FIGURES')
replaced = 0
for fragment, img_path, label in UPDATES:
    idx = find_idx(prs, fragment)
    if idx is None:
        print(f'  WARNING: could not find slide for "{fragment}" ({label})')
        continue
    print(f'  {label}  idx={idx}  "{fragment[:50]}"')
    if replace_figure(prs.slides[idx], img_path, prs, label):
        replaced += 1

print(f'\nReplaced {replaced}/{len(UPDATES)} figures')
print(f'Final: {len(prs.slides)} slides')
for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved -> {out}')

print('\nPOST-BUILD VERIFICATION')
for out in OUTS:
    if os.path.exists(out):
        verify_zip(out)
