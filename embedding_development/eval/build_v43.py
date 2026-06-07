#!/usr/bin/env python3
"""
build_v43.py  —  All-in-one from v41: freq trace fix + complete E07 cue-zone rebuild.

Base: user's Desktop ultimate_presentation_v41.pptx  (avoids orphaned-part accumulation)
v41 (93) → v43 (93 slides):
  S36: freq_trace_example.png (opaque legend fix)
  S68: e07_joint_tuning.png  → e07_cue_zone_tuning.png
       "E07 — cue-specific position tuning (top 4 sessions)"
       → "E07 in the cue zone: Cue 1 vs Cue 2 position tuning"
  S69: e07_position_cue_tuning.png → e07_cue_zone_joint_gpv.png
       title → "Cue zone joint GPV: interaction ≈0 — still additive"
  S70: "Signal routes through speed..." → replaced with e07_cue_zone_attribution.png
       title → "Head angle and speed capture cue identity in the cue zone"
"""
import os, zipfile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.parts.presentation import PresentationPart


@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
mdir = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')
adir = os.path.join(root, 'outputs', 'ablation_vs_attribution')

SRC  = '/mnt/c/Users/amits/Desktop/ultimate_presentation_v41.pptx'
OUTS = [
    os.path.join(root, 'outputs', 'ultimate_presentation_v43.pptx'),
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v43.pptx',
]

IMG_FREQ = os.path.join(mdir, 'freq_trace_example.png')
IMG_68   = os.path.join(adir, 'e07_cue_zone_tuning.png')
IMG_69   = os.path.join(adir, 'e07_cue_zone_joint_gpv.png')
IMG_70   = os.path.join(adir, 'e07_cue_zone_attribution.png')
TARGET_DPI = 200

TITLE_68 = 'E07 in the cue zone: Cue 1 vs Cue 2 position tuning'
TITLE_69 = 'Cue zone joint GPV: interaction ≈0 — still additive'
TITLE_70 = 'Head angle and speed capture cue identity in the cue zone'

C_BG    = RGBColor(0xFF, 0xFF, 0xFF)
C_TITLE = RGBColor(0x1A, 0x23, 0x3A)


def verify_figure(path, label):
    if not os.path.exists(path):
        print(f'  ✗ MISSING: {label}'); return False
    with Image.open(path) as im:
        w, h = im.size
    print(f'  ✓ {label}: {w}\xd7{h}px = {w/TARGET_DPI:.2f}"\xd7{h/TARGET_DPI:.2f}"')
    return True


def verify_zip(path):
    with zipfile.ZipFile(path) as z:
        names  = z.namelist()
        dupes  = [n for n in set(names) if names.count(n) > 1]
        slides = [n for n in names if 'slides/slide' in n
                  and n.endswith('.xml') and 'rels' not in n]
    ok = len(dupes) == 0
    print(f'  {chr(10003) if ok else chr(10007)+" DUPES: "+str(dupes)} {len(slides)} slides, {len(dupes)} dupes')
    return ok


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def replace_image(slide, new_path, left, top, w, h):
    for sh in list(slide.shapes):
        if sh.shape_type == 13:
            sh._element.getparent().remove(sh._element)
            break
    slide.shapes.add_picture(new_path,
        int(Inches(left)), int(Inches(top)),
        int(Inches(w)),    int(Inches(h)))


def patch_title(slide, new_text):
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            if para.runs:
                para.runs[0].text = new_text
                for r in list(para.runs)[1:]:
                    r.text = ''
                return True
    return False


def move_slide(prs, from_idx, to_idx):
    xml = prs.slides._sldIdLst
    els = list(xml)
    el  = els.pop(from_idx)
    els.insert(to_idx, el)
    for c in list(xml): xml.remove(c)
    for c in els:       xml.append(c)


# ── Pre-build verification ─────────────────────────────────────────────────────
print('PRE-BUILD VERIFICATION')
ok  = verify_figure(IMG_FREQ, 'freq_trace_example')
ok &= verify_figure(IMG_68,   'e07_cue_zone_tuning')
ok &= verify_figure(IMG_69,   'e07_cue_zone_joint_gpv')
ok &= verify_figure(IMG_70,   'e07_cue_zone_attribution')
for title, label in [(TITLE_68,'S68'),(TITLE_69,'S69'),(TITLE_70,'S70')]:
    assert len(title) <= 75, f'{label} title too long: {len(title)} chars'
    print(f'  ✓ {label} title: {len(title)} chars')
if not os.path.exists(SRC):
    raise SystemExit(f'Source not found: {SRC}')
print(f'  ✓ source: {os.path.basename(SRC)}\n')
if not ok:
    raise SystemExit('Verification failed')

prs = Presentation(SRC)
print(f'Opened: {len(prs.slides)} slides')

# ── S36: fix freq trace legend ────────────────────────────────────────────────
idx_freq = None
for probe in range(32, 40):
    sl = prs.slides[probe]
    for sh in sl.shapes:
        if sh.shape_type == 13 and len(sh.image.blob) > 280000:
            idx_freq = probe; break
    if idx_freq is not None: break
if idx_freq is None:
    raise SystemExit('Could not find freq trace slide')
replace_image(prs.slides[idx_freq], IMG_FREQ, 0.25, 0.62, 9.50, 4.20)
print(f'  S{idx_freq+1}: freq_trace_example replaced')

# ── S68: replace per-session tuning → cue-zone tuning ─────────────────────────
idx68 = find_idx(prs, 'joint position')
if idx68 is None:
    idx68 = find_idx(prs, 'E07: joint')
if idx68 is None:
    raise SystemExit('Could not find S68')
replace_image(prs.slides[idx68], IMG_68, 0.25, 0.68, 9.50, 4.20)
patch_title(prs.slides[idx68], TITLE_68)
print(f'  S{idx68+1}: replaced with cue_zone_tuning')

# ── S69: replace empty averaged curves → cue-zone joint GPV ───────────────────
idx69 = find_idx(prs, 'position tuning averaged')
if idx69 is None:
    idx69 = find_idx(prs, 'averaged across sessions')
if idx69 is None:
    raise SystemExit('Could not find S69')
replace_image(prs.slides[idx69], IMG_69, 0.25, 0.62, 9.50, 4.20)
patch_title(prs.slides[idx69], TITLE_69)
print(f'  S{idx69+1}: replaced with cue_zone_joint_gpv')

# ── S70: replace "Signal routes through speed" with cue-zone attribution ───────
idx70 = find_idx(prs, 'Signal routes through speed')
if idx70 is None:
    raise SystemExit('Could not find S70 (Signal routes through speed)')
replace_image(prs.slides[idx70], IMG_70, 0.25, 0.62, 9.50, 4.20)
patch_title(prs.slides[idx70], TITLE_70)
print(f'  S{idx70+1}: replaced with cue_zone_attribution')

# ── Context ───────────────────────────────────────────────────────────────────
print('\nE07 case study context:')
for i in range(idx68 - 1, min(len(prs.slides), idx68 + 5)):
    sl    = prs.slides[i]
    texts = [sh.text_frame.text.strip()[:65] for sh in sl.shapes
             if sh.has_text_frame and sh.text_frame.text.strip()]
    imgs  = any(sh.shape_type == 13 for sh in sl.shapes)
    print(f'  S{i+1:02d}: {texts[0] if texts else "(empty)"}{"  [IMG]" if imgs else ""}')

print(f'\nFinal: {len(prs.slides)} slides')
for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved → {out}')

print('\nPOST-BUILD VERIFICATION')
for out in OUTS:
    if os.path.exists(out):
        verify_zip(out)
