#!/usr/bin/env python3
"""
build_v23.py  —  Six structural fixes from user feedback, v22 (85) → v23 (85 slides).

Changes:
  1. Remove S46 honest_attribution_example from attribution section
  2. Fix S48 attribution_consistency_merged: re-insert at correct aspect ratio
     (9.50"×4.20" figure was inserted at 5.30"×4.95" — wrong)
  3. Fix S50 cross_ensemble_prediction: re-insert at correct aspect ratio
     (6.00"×4.50" figure was inserted at 9.50"×3.89" — wrong)
  4. Head angle section:
       a. Strip embedded image from section header → clean divider
       b. Insert new head_angle_attribution_summary slide as first content slide
  5. Move honest_attribution_example to position section (before position ablation)
     with updated title; reconciles "position IS predictive" (ablation) with
     "head angle outcompetes it in the full model" (attribution)
  6. E07/E23 section redesign:
       a. Drop old intro image slide (old image from cleanup_v3, no narrative)
       b. Update section header title to state the finding

Net: 85 slides (add 2 summary/honest slides, remove 2 old slides).
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
mdir = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v22.pptx')
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v23.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation_v23.pptx'),
]

SLIDE_W  = Inches(10.0)
SLIDE_H  = Inches(5.625)
C_BG     = RGBColor(0xFF, 0xFF, 0xFF)
C_TITLE  = RGBColor(0x1A, 0x23, 0x3A)
C_SUB    = RGBColor(0x55, 0x55, 0x55)


# ── Helpers ───────────────────────────────────────────────────────────────────

def move_slide(prs, from_idx, to_idx):
    xml = prs.slides._sldIdLst
    els = list(xml)
    el  = els.pop(from_idx)
    els.insert(to_idx, el)
    for c in list(xml): xml.remove(c)
    for c in els:        xml.append(c)


def remove_slide(prs, idx):
    slide      = prs.slides[idx]
    slide_part = slide.part
    prs_part   = prs.slides.part
    rId = None
    for rel_id, rel in prs_part.rels.items():
        try:
            if rel.target_part is slide_part:
                rId = rel_id; break
        except Exception:
            pass
    if rId:
        prs_part.rels.pop(rId)
    xml = prs.slides._sldIdLst
    xml.remove(list(xml)[idx])


def find_idx(prs, fragment):
    """Return index of first slide whose text contains fragment (case-insensitive)."""
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def fix_picture(slide, img_path, left_in, top_in, width_in, height_in):
    """Replace the first picture on a slide with img_path at given dimensions."""
    for sh in list(slide.shapes):
        if sh.shape_type == 13:
            sh._element.getparent().remove(sh._element)
            break
    slide.shapes.add_picture(img_path,
        int(Inches(left_in)), int(Inches(top_in)),
        int(Inches(width_in)), int(Inches(height_in)))


def strip_pictures(slide):
    """Remove all embedded pictures from a slide (make it a clean text slide)."""
    removed = 0
    for sh in list(slide.shapes):
        if sh.shape_type == 13:
            sh._element.getparent().remove(sh._element)
            removed += 1
    return removed


def patch_text(slide, old, new):
    found = False
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            for run in para.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    found = True
    return found


def add_figure_slide(prs, title_text, img_path,
                     left_in=0.25, top_in=0.60, width_in=9.50, height_in=4.80,
                     subtitle_text=None):
    """Add a blank slide with a bold title and a figure image."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill  = slide.background.fill; fill.solid(); fill.fore_color.rgb = C_BG

    tb = slide.shapes.add_textbox(Inches(0.20), Inches(0.06), Inches(9.60), Inches(0.52))
    p  = tb.text_frame.paragraphs[0]
    r  = p.add_run(); r.text = title_text
    r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = C_TITLE

    if subtitle_text:
        tb2 = slide.shapes.add_textbox(Inches(0.20), Inches(0.56), Inches(9.60), Inches(0.22))
        p2  = tb2.text_frame.paragraphs[0]
        r2  = p2.add_run(); r2.text = subtitle_text
        r2.font.size = Pt(11); r2.font.color.rgb = C_SUB
        top_in = 0.78

    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path,
            int(Inches(left_in)), int(Inches(top_in)),
            int(Inches(width_in)), int(Inches(height_in)))
    else:
        print(f'  ⚠  Image not found: {img_path}')

    return len(prs.slides) - 1


# ── Load ──────────────────────────────────────────────────────────────────────
prs = Presentation(SRC)
print(f'Opened v22: {len(prs.slides)} slides')


# ══ STEP 1: Add new slides to end (must happen before any removes) ════════════

# New slide A: head_angle_attribution_summary
new_ha_idx = add_figure_slide(
    prs,
    title_text='Head Angle Is the Top-Attributed Feature Across All Metrics',
    img_path=os.path.join(mdir, 'head_angle_attribution_summary.png'),
    left_in=0.25, top_in=0.60, width_in=9.50, height_in=3.50,
)
print(f'  Added head_angle_attribution_summary at idx {new_ha_idx}')

# New slide B: honest_attribution_example (for position section)
new_honest_idx = add_figure_slide(
    prs,
    title_text='Position Appears Primary by Correlation — Head Angle Outcompetes It in the Full Model',
    img_path=os.path.join(mdir, 'honest_attribution_example.png'),
    left_in=0.25, top_in=0.62, width_in=9.50, height_in=4.80,
)
print(f'  Added honest_attribution_example at idx {new_honest_idx}')
# → 87 slides


# ══ STEP 2: Remove old slides ════════════════════════════════════════════════

# 2a. Remove old honest_attribution from attribution section (S46 in v22, idx 45)
old_honest_idx = find_idx(prs, 'Attribution Is More Honest Than Correlation')
if old_honest_idx is not None:
    remove_slide(prs, old_honest_idx)
    print(f'  Removed old honest_attribution from idx {old_honest_idx}')
    # new slides at end shifted down by 1
    new_ha_idx     -= 1
    new_honest_idx -= 1
else:
    print('  ⚠  Could not find old honest_attribution slide')
# → 86 slides

# 2b. Remove E07/E23 intro image slide
e07_intro_idx = find_idx(prs, 'Case Study: E07-Cue Visible')
if e07_intro_idx is not None:
    remove_slide(prs, e07_intro_idx)
    print(f'  Removed E07/E23 intro image from idx {e07_intro_idx}')
    if new_ha_idx     > e07_intro_idx: new_ha_idx     -= 1
    if new_honest_idx > e07_intro_idx: new_honest_idx -= 1
else:
    print('  ⚠  Could not find E07/E23 intro slide')
# → 85 slides


# ══ STEP 3: Fix stretched pictures ═══════════════════════════════════════════

# 3a. attribution_consistency_merged.png: 9.50"×4.20" → insert at 9.40"×4.16"
ac_idx = find_idx(prs, 'Attribution Profile Consistency')
if ac_idx is not None:
    fix_picture(prs.slides[ac_idx],
                os.path.join(mdir, 'attribution_consistency_merged.png'),
                left_in=0.30, top_in=0.62, width_in=9.40, height_in=4.16)
    print(f'  Fixed attribution_consistency_merged on S{ac_idx+1}')
else:
    print('  ⚠  Could not find Attribution Profile Consistency slide')

# 3b. cross_ensemble_prediction.png: 6.00"×4.50" → centred at (2.00", 0.60")
ce_idx = find_idx(prs, 'Cross-Ensemble Neural Prediction')
if ce_idx is not None:
    fix_picture(prs.slides[ce_idx],
                os.path.join(mdir, 'cross_ensemble_prediction.png'),
                left_in=2.00, top_in=0.60, width_in=6.00, height_in=4.50)
    print(f'  Fixed cross_ensemble_prediction on S{ce_idx+1}')
else:
    print('  ⚠  Could not find Cross-Ensemble Neural Prediction slide')


# ══ STEP 4: Head angle section header ════════════════════════════════════════

ha_header_idx = find_idx(prs, 'Head Angle Tuning Curves')
if ha_header_idx is not None:
    n_removed = strip_pictures(prs.slides[ha_header_idx])
    print(f'  Stripped {n_removed} picture(s) from head angle section header S{ha_header_idx+1}')
    # Update the slide number label '51' if present (cosmetic)
    patch_text(prs.slides[ha_header_idx], '51', '')
else:
    print('  ⚠  Could not find Head Angle Tuning Curves slide')


# ══ STEP 5: Move new slides to target positions ════════════════════════════════

# 5a. Move head_angle_attribution_summary to just after head angle section header
ha_header_idx = find_idx(prs, 'Head Angle Tuning Curves')   # re-find after removes
if ha_header_idx is not None:
    target = ha_header_idx + 1
    move_slide(prs, new_ha_idx, target)
    print(f'  Moved head_angle_attribution_summary to idx {target} (S{target+1})')
    # Update honest_attribution index if it was after target
    if new_honest_idx >= target and new_honest_idx != new_ha_idx:
        new_honest_idx += 1
    new_ha_idx = target
else:
    print('  ⚠  Could not find head angle header for placement')

# 5b. Move honest_attribution_example to just before position ablation
pos_ablation_idx = find_idx(prs, 'Position-Only MLP Recovers')
if pos_ablation_idx is None:
    pos_ablation_idx = find_idx(prs, 'Position IS Predictive')
if pos_ablation_idx is not None:
    # re-find honest idx (may have shifted after 5a)
    new_honest_idx = find_idx(prs, 'Position Appears Primary by Correlation')
    if new_honest_idx is not None:
        move_slide(prs, new_honest_idx, pos_ablation_idx)
        print(f'  Moved honest_attribution to idx {pos_ablation_idx} (S{pos_ablation_idx+1}), before position ablation')
    else:
        print('  ⚠  Could not find new honest_attribution slide')
else:
    print('  ⚠  Could not find position ablation slide')


# ══ STEP 6: E07/E23 section header narrative update ══════════════════════════

e07_header_idx = find_idx(prs, 'Case Studies: Existing Validation')
if e07_header_idx is None:
    e07_header_idx = find_idx(prs, 'E07 and E23 attribution in depth')
if e07_header_idx is not None:
    patch_text(prs.slides[e07_header_idx],
               'Case Studies: Existing Validation',
               'E07: Speed-Dominated  ·  E23: Head Angle-Dominated')
    patch_text(prs.slides[e07_header_idx],
               'E07 and E23 attribution in depth',
               'Attribution profiles validated against independent neural measures')
    print(f'  Updated E07/E23 section header on S{e07_header_idx+1}')
else:
    print('  ⚠  Could not find E07/E23 section header')


# ══ Report ════════════════════════════════════════════════════════════════════
print(f'\nFinal slide count: {len(prs.slides)}')

# Print key slides for verification
key_fragments = [
    'Attribution Is More Honest',           # should be GONE
    'Head Angle Is the Top-Attributed',     # new S51
    'Position Appears Primary by Corr',     # new ~S59
    'Position-Only MLP',                    # position ablation
    'E07: Speed-Dominated',                 # updated E07/E23 header
    'Case Study: E07-Cue',                  # should be GONE
    'Attribution Profile Consistency',      # fixed S48
    'Cross-Ensemble Neural Prediction',     # fixed S50
]
print('\nKey slide check:')
for frag in key_fragments:
    idx = find_idx(prs, frag)
    status = f'S{idx+1} (idx {idx})' if idx is not None else '✗ NOT FOUND'
    print(f'  {frag[:50]:<50} → {status}')

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'\nSaved → {out}')
