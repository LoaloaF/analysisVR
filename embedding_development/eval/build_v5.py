#!/usr/bin/env python3
"""
build_v5.py  —  Reorder v4 slides to the new narrative structure and save as v5.

New flow (section → 0-based indices from v4):
  Intro            : 0–10
  Linear           : 11, 12
  MLP              : 13, 14, 15
  TempConv + R²    : 31, 32, 16, 17
  Example preds    : 43, 44, 45
  Attribution      : 20, 21, 22, 23, 24, 25, 26, 27, 39, 40, 41, 42
  Consistency      : 18, 19
  Frequency        : 30
  Case studies     : 33, 34, 35, 36, 37, 38, 28, 29
"""

import os
from pptx import Presentation

SRC  = '/mnt/c/Users/amits/Desktop/ultimate_presentation_v4.pptx'
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v5.pptx',
    os.path.join(os.path.dirname(os.path.abspath(__file__)), '..',
                 'outputs', 'ultimate_presentation_v5.pptx'),
]

NEW_ORDER = [
    # ── Intro ──────────────────────────────────────────────────────────────────
    0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10,
    # ── Linear model + R² ──────────────────────────────────────────────────────
    11, 12,
    # ── MLP + R² ───────────────────────────────────────────────────────────────
    13, 14, 15,
    # ── TempConv intro + grand-mean R² comparison ──────────────────────────────
    31, 32, 16, 17,
    # ── Example prediction slides ──────────────────────────────────────────────
    43, 44, 45,
    # ── Feature attribution: intro, MLP results, cross-model, validation ───────
    20, 21, 22, 23, 24, 25, 26, 27, 39, 40, 41, 42,
    # ── Embedding consistency ──────────────────────────────────────────────────
    18, 19,
    # ── Frequency analysis ─────────────────────────────────────────────────────
    30,
    # ── Case studies ───────────────────────────────────────────────────────────
    33, 34, 35, 36, 37, 38, 28, 29,
]

assert len(NEW_ORDER) == 46, f"Expected 46, got {len(NEW_ORDER)}"
assert len(set(NEW_ORDER)) == 46, "Duplicate indices in NEW_ORDER"

prs = Presentation(SRC)
assert len(prs.slides) == 46, f"Expected 46 slides in v4, got {len(prs.slides)}"

xml_slides = prs.slides._sldIdLst
all_els    = list(xml_slides)

for el in list(xml_slides):
    xml_slides.remove(el)
for idx in NEW_ORDER:
    xml_slides.append(all_els[idx])

print(f"Reordered {len(prs.slides)} slides")

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f"Saved → {out}")

print("Done.")
