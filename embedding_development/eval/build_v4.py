#!/usr/bin/env python3
"""
build_v4.py

Open ultimate_presentation_v3.pptx (user's edited version), replace all
stale figure images with freshly regenerated Arial-font versions, save as v4.

Two categories of replacements:
  1. Figure slides (white solid bg): one picture per slide — match by title prefix.
  2. Story slides with embedded matplotlib images (11, 29, 34, 44-46): replace
     each picture's blob directly.
"""

import os, sys
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.enum.dml import MSO_FILL

SRC  = '/mnt/c/Users/amits/Desktop/ultimate_presentation_v3.pptx'
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v4.pptx',
    os.path.join(os.path.dirname(os.path.abspath(__file__)), '..',
                 'outputs', 'ultimate_presentation_v4.pptx'),
]

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, "..")

# ── Figure file lookup: search all output dirs ──────────────────────────────────
FIGURE_DIRS = [
    os.path.join(root, 'outputs', 'cebra_comparison'),
    os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed'),
    os.path.join(root, 'outputs', 'mlps', 'ml_vs_naive'),
    os.path.join(root, 'outputs', 'cebra_eval', 'ensembles'),
    os.path.join(root, 'outputs', 'temporal_advantage'),
    os.path.join(root, 'outputs', 'mlps', 'head_angle_tuning'),
    os.path.join(root, 'outputs', 'trial_traces_comparison'),
    os.path.join(root, 'outputs', 'supervisor_figures'),
    os.path.join(root, 'outputs', 'mlps', 'spikes_multiseed'),
]

figs = {}
for d in FIGURE_DIRS:
    if os.path.isdir(d):
        for fn in os.listdir(d):
            if fn.endswith('.png'):
                figs[fn] = os.path.join(d, fn)

print(f"Found {len(figs)} PNG files")

# ── Title → filename map for figure slides ─────────────────────────────────────
TITLE_TO_FILE = {
    'linear baseline':             'r2_bar_linear.png',
    'mlp baseline':                'r2_bar_mlp.png',
    'grand mean r':                'r2_grand_mean_bars.png',
    'valid (r':                    'two_thresholds_scatter.png',
    'mlp cross-seed':              'r2_consistency_bar.png',
    'embedding consistency':       'embedding_consistency_comparison.png',
    'global permutation variance': 'gpv_group_ensemble_heatmap.png',
    'global pv vs':                'global_vs_cond_pv_scatter.png',
    'mean |ig|':                   'ig_per_ensemble_heatmap.png',
    'group-level attribution':     'group_attribution_comparison.png',
    'attribution agreement':       'attribution_agreement_scatter.png',
    'case studies: e07':           'case_studies_e07_e23.png',
    'mlp fails on fast':           'trend_noise_comparison.png',
    'top pairs: head angle':       'head_angle_scatter_2x2.png',
    'e18 tuning curve stability':  'head_angle_stability.png',
    'e18 tuning shape variety':    'head_angle_tuning_6panel.png',
    'position tuning curves':      'position_tuning.png',
    'part 1':                      'ml_vs_naive_scatter.png',
    'part 2':                      'nonlinearity_advantage.png',
    'part 3':                      'mlp_vs_linear_r2.png',
    'part 4':                      'ablation_proof.png',
}


def replace_first_image(slide, new_path):
    """Replace the blob of the first picture shape in slide with new_path bytes."""
    for s in slide.shapes:
        if hasattr(s, '_pic'):
            blip = s._pic.find('.//' + qn('a:blip'))
            if blip is not None:
                rid = blip.get(qn('r:embed'))
                if rid:
                    part = slide.part.related_part(rid)
                    with open(new_path, 'rb') as f:
                        part._blob = f.read()
                    return True
    return False


def replace_nth_image(slide, n, new_path):
    """Replace the n-th (0-indexed) picture's blob in slide."""
    count = 0
    for s in slide.shapes:
        if hasattr(s, '_pic'):
            blip = s._pic.find('.//' + qn('a:blip'))
            if blip is not None:
                rid = blip.get(qn('r:embed'))
                if rid:
                    if count == n:
                        part = slide.part.related_part(rid)
                        with open(new_path, 'rb') as f:
                            part._blob = f.read()
                        return True
                    count += 1
    return False


prs = Presentation(SRC)
print(f"Loaded v3: {len(prs.slides)} slides")

n_replaced = 0

for i, slide in enumerate(prs.slides):
    slide_num = i + 1

    # ── Figure slides (white solid bg) ──────────────────────────────────────
    if slide.background.fill.type == MSO_FILL.SOLID:
        # Get title text
        title = ''
        for s in slide.shapes:
            if s.has_text_frame and s.text.strip():
                title = s.text.strip().lower()[:60]
                break

        # Match to file
        fname = None
        for prefix, fn in TITLE_TO_FILE.items():
            if title.startswith(prefix.lower()):
                fname = fn
                break

        if fname and fname in figs:
            if replace_first_image(slide, figs[fname]):
                print(f"  [{slide_num:02d}] FIG replaced: {fname}")
                n_replaced += 1
            else:
                print(f"  [{slide_num:02d}] FIG no picture found for: {fname}")
        elif fname:
            print(f"  [{slide_num:02d}] FIG file not found: {fname}")
        continue

    # ── Story slides with specific embedded images ──────────────────────────
    if slide_num == 11:
        # Two variance scatter plots: img0=MSE, img1=R²
        if 'variance_vs_mse.png' in figs and replace_nth_image(slide, 0, figs['variance_vs_mse.png']):
            print(f"  [11] STORY replaced img0: variance_vs_mse.png")
            n_replaced += 1
        if 'variance_vs_r2.png' in figs and replace_nth_image(slide, 1, figs['variance_vs_r2.png']):
            print(f"  [11] STORY replaced img1: variance_vs_r2.png")
            n_replaced += 1

    elif slide_num == 29:
        if 'fig1_data_distributions.png' in figs and replace_first_image(slide, figs['fig1_data_distributions.png']):
            print(f"  [29] STORY replaced: fig1_data_distributions.png")
            n_replaced += 1

    elif slide_num == 34:
        if 'tuning_curves_by_category.png' in figs and replace_first_image(slide, figs['tuning_curves_by_category.png']):
            print(f"  [34] STORY replaced: tuning_curves_by_category.png")
            n_replaced += 1

    elif slide_num == 44:
        if 'overlay_S01_E03.png' in figs and replace_first_image(slide, figs['overlay_S01_E03.png']):
            print(f"  [44] STORY replaced: overlay_S01_E03.png")
            n_replaced += 1

    elif slide_num == 45:
        if 'overlay_S07_E04.png' in figs and replace_first_image(slide, figs['overlay_S07_E04.png']):
            print(f"  [45] STORY replaced: overlay_S07_E04.png")
            n_replaced += 1

    elif slide_num == 46:
        if 'overlay_S25_E05.png' in figs and replace_first_image(slide, figs['overlay_S25_E05.png']):
            print(f"  [46] STORY replaced: overlay_S25_E05.png")
            n_replaced += 1

print(f"\n{n_replaced} images replaced total")

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f"Saved ({len(prs.slides)} slides) → {out}")

print("Done.")
