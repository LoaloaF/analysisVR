#!/usr/bin/env python3
"""
build_v31.py  —  Shorter slide titles + consistent scatter axis bounds.

v30 (86) → v31 (86 slides, same count):
  S62-S64 titles shortened; all three figures regenerated with cleaner panel titles
  and consistent shared x/y limits on the S64 scatter.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.parts.presentation import PresentationPart


# Fix python-pptx bug: _next_slide_partname uses len(sldIdLst)+1 which conflicts with
# existing slide parts after removals. Use package.next_partname instead.
@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
adir = os.path.join(root, 'outputs', 'ablation_vs_attribution')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v30.pptx')
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v31.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation_v31.pptx'),
]

C_BG    = RGBColor(0xFF, 0xFF, 0xFF)
C_TITLE = RGBColor(0x1A, 0x23, 0x3A)


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def remove_slide(prs, idx):
    slide = prs.slides[idx]; slide_part = slide.part; prs_part = prs.slides.part
    rId = None
    for rel_id, rel in prs_part.rels.items():
        try:
            if rel.target_part is slide_part: rId = rel_id; break
        except Exception: pass
    if rId: prs_part.rels.pop(rId)
    xml = prs.slides._sldIdLst; xml.remove(list(xml)[idx])
    # Purge orphaned part from OPC store so its filename is freed for new slides
    try:
        prs_part.package._PartStore._parts.pop(slide_part.partname, None)
    except Exception:
        pass


def move_slide(prs, from_idx, to_idx):
    xml = prs.slides._sldIdLst
    els = list(xml); el = els.pop(from_idx); els.insert(to_idx, el)
    for c in list(xml): xml.remove(c)
    for c in els:        xml.append(c)


def add_figure_slide(prs, title_text, img_path, fig_top=0.62, fig_h=4.20):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill  = slide.background.fill; fill.solid(); fill.fore_color.rgb = C_BG
    tb = slide.shapes.add_textbox(Inches(0.20), Inches(0.06), Inches(9.60), Inches(0.52))
    p  = tb.text_frame.paragraphs[0]
    r  = p.add_run(); r.text = title_text
    r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = C_TITLE
    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path,
            int(Inches(0.25)), int(Inches(fig_top)),
            int(Inches(9.50)), int(Inches(fig_h)))
    else:
        print(f'  ⚠  Missing: {img_path}')
    return len(prs.slides) - 1


prs = Presentation(SRC)
print(f'Opened v30: {len(prs.slides)} slides')

# Step 1: remove the three old slides (search by title fragment, always remove from top)
OLD_FRAGMENTS = [
    "Where the Supervisor Has a Point",   # S62 from v27
    "Bucket 1 — Our Model Detects",       # S63 from v27
    "Bucket 2 — We Don't Miss",           # S64 from v27
]
for frag in OLD_FRAGMENTS:
    idx = find_idx(prs, frag)
    if idx is not None:
        remove_slide(prs, idx)
        print(f'  Removed "{frag}" at idx {idx}')
    else:
        print(f'  ⚠  Not found: "{frag}"')
print(f'After removals: {len(prs.slides)} slides')

# Step 2: find insertion point (after Case Studies header)
header_idx = find_idx(prs, 'Case Studies')
if header_idx is None:
    header_idx = len(prs.slides) - 1
    print(f'  ⚠  "Case Studies" not found; appending at end')
target = header_idx + 1
print(f'  Inserting after "Case Studies" (S{header_idx+1}) → target idx {target}')

# Step 3: add new slides at end, then move into position
idx_s1 = add_figure_slide(
    prs,
    'Supervisor\'s variable has a neural signal',
    os.path.join(adir, 'e07_e23_bucket_signal.png'))
print(f'  Added bucket_signal at idx {idx_s1}')

idx_s2 = add_figure_slide(
    prs,
    'Bucket 1: model attributes to their variable (GPV/R² ≥ 10%)',
    os.path.join(adir, 'e07_e23_bucket_gpv.png'))
print(f'  Added bucket_gpv at idx {idx_s2}')

idx_s3 = add_figure_slide(
    prs,
    'Bucket 2: a co-varying feature captures the signal',
    os.path.join(adir, 'e07_e23_bucket_covariation.png'))
print(f'  Added bucket_covariation at idx {idx_s3}')

# Move slides into position
move_slide(prs, idx_s1, target)
print(f'  Moved bucket_signal → idx {target}')

idx_s2 = find_idx(prs, 'Bucket 1: model attributes')
move_slide(prs, idx_s2, target + 1)
print(f'  Moved bucket_gpv → idx {target + 1}')

idx_s3 = find_idx(prs, 'Bucket 2: a co-varying')
move_slide(prs, idx_s3, target + 2)
print(f'  Moved bucket_covariation → idx {target + 2}')

# ── Verify ────────────────────────────────────────────────────────────────────
print(f'\nFinal slide count: {len(prs.slides)}')
header_idx = find_idx(prs, 'Case Studies')
print('\nCase Studies section:')
for i in range(header_idx, min(header_idx + 5, len(prs.slides))):
    sl = prs.slides[i]
    texts = [sh.text_frame.text.strip()[:80]
             for sh in sl.shapes if sh.has_text_frame]
    title = next((t for t in texts if t), '(empty)')
    print(f'  S{i+1:02d}: {title}')

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'\nSaved → {out}')
