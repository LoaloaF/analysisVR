#!/usr/bin/env python3
"""
build_v44.py  —  Update S69 and S70 with cue-zone separation/features figures.

Base: outputs/ultimate_presentation_v43.pptx
v43 (93) → v44 (93 slides):
  S69: e07_cue_zone_joint_gpv.png → e07_cue_zone_separation.png
       title → "S20: cue-zone activity and MLP predictions — Cue 1 vs Cue 2"
  S70: e07_cue_zone_attribution.png → e07_cue_zone_features.png
       title → "In the cue zone: head angle and speed co-vary with cue identity"
"""
import os, zipfile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from pptx.parts.presentation import PresentationPart


@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
adir = os.path.join(root, 'outputs', 'ablation_vs_attribution')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v43.pptx')
OUTS = [
    os.path.join(root, 'outputs', 'ultimate_presentation_v44.pptx'),
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v44.pptx',
]

IMG_69 = os.path.join(adir, 'e07_cue_zone_separation.png')
IMG_70 = os.path.join(adir, 'e07_cue_zone_features.png')
TARGET_DPI = 200

TITLE_69 = 'S20: cue-zone activity and MLP predictions — Cue 1 vs Cue 2'
TITLE_70 = 'In the cue zone: head angle and speed co-vary with cue identity'


def verify_figure(path, label):
    if not os.path.exists(path):
        print(f'  ✗ MISSING: {label}'); return False
    with Image.open(path) as im:
        w, h = im.size
    print(f'  ✓ {label}: {w}x{h}px = {w/TARGET_DPI:.2f}"x{h/TARGET_DPI:.2f}"')
    return True


def verify_zip(path):
    with zipfile.ZipFile(path) as z:
        names  = z.namelist()
        dupes  = [n for n in set(names) if names.count(n) > 1]
        slides = [n for n in names if 'slides/slide' in n
                  and n.endswith('.xml') and 'rels' not in n]
    ok = len(dupes) == 0
    mark = chr(10003) if ok else chr(10007) + ' DUPES: ' + str(dupes)
    print(f'  {mark} {len(slides)} slides, {len(dupes)} dupes')
    return ok


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def replace_image(slide, new_path, left, top, w, h):
    for sh in list(slide.shapes):
        if sh.shape_type == 13:
            sh._element.getparent().remove(sh._element)
            break
    slide.shapes.add_picture(new_path,
        int(Inches(left)), int(Inches(top)),
        int(Inches(w)),    int(Inches(h)))


def patch_title(slide, new_text):
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            if para.runs:
                para.runs[0].text = new_text
                for r in list(para.runs)[1:]:
                    r.text = ''
                return True
    return False


print('PRE-BUILD VERIFICATION')
ok  = verify_figure(IMG_69, 'e07_cue_zone_separation')
ok &= verify_figure(IMG_70, 'e07_cue_zone_features')
for title, label in [(TITLE_69, 'S69'), (TITLE_70, 'S70')]:
    assert len(title) <= 75, f'{label} title too long: {len(title)}'
    print(f'  ✓ {label} title: {len(title)} chars')
if not os.path.exists(SRC):
    raise SystemExit(f'Source not found: {SRC}')
print(f'  ✓ source: {os.path.basename(SRC)}\n')
if not ok:
    raise SystemExit('Verification failed')

prs = Presentation(SRC)
print(f'Opened: {len(prs.slides)} slides')

# S69
idx69 = find_idx(prs, 'Cue zone joint GPV')
if idx69 is None:
    idx69 = find_idx(prs, 'interaction')
if idx69 is None:
    raise SystemExit('Could not find S69')
replace_image(prs.slides[idx69], IMG_69, 0.25, 0.62, 9.50, 4.20)
patch_title(prs.slides[idx69], TITLE_69)
print(f'  S{idx69+1}: cue_zone_separation inserted')

# S70
idx70 = find_idx(prs, 'Head angle and speed capture cue identity')
if idx70 is None:
    idx70 = find_idx(prs, 'capture cue identity')
if idx70 is None:
    raise SystemExit('Could not find S70')
replace_image(prs.slides[idx70], IMG_70, 0.25, 0.62, 9.50, 4.20)
patch_title(prs.slides[idx70], TITLE_70)
print(f'  S{idx70+1}: cue_zone_features inserted')

# Context
print('\nE07 case study:')
for i in range(idx69 - 2, min(len(prs.slides), idx69 + 5)):
    sl    = prs.slides[i]
    texts = [sh.text_frame.text.strip()[:65] for sh in sl.shapes
             if sh.has_text_frame and sh.text_frame.text.strip()]
    imgs  = any(sh.shape_type == 13 for sh in sl.shapes)
    print(f'  S{i+1:02d}: {texts[0] if texts else "(empty)"}{"  [IMG]" if imgs else ""}')

print(f'\nFinal: {len(prs.slides)} slides')
for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved → {out}')

print('\nPOST-BUILD VERIFICATION')
for out in OUTS:
    if os.path.exists(out):
        verify_zip(out)
