#!/usr/bin/env python3
"""
build_v33.py  —  Replace S51 and S60 figures; S50 untouched.

v32 (87) → v33 (87 slides, same count):
  S51: head_angle_attribution_summary.png  — removed '#1 by GPV and IG' annotation
  S60: position_ablation.png               — Full bar colour grey (not green); cache added
  S50: Head Angle Analysis title slide     — NOT modified (retained as-is)
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.parts.presentation import PresentationPart


@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
mdir = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')

SRC  = os.path.join(root, 'outputs', 'ultimate_presentation_v32.pptx')
OUTS = [
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v33.pptx',
    os.path.join(root, 'outputs', 'ultimate_presentation_v33.pptx'),
]


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def replace_picture(slide, new_img, left_in, top_in, w_in, h_in):
    for sh in list(slide.shapes):
        if sh.shape_type == 13:
            sh._element.getparent().remove(sh._element)
            break
    slide.shapes.add_picture(new_img,
        int(Inches(left_in)), int(Inches(top_in)),
        int(Inches(w_in)),    int(Inches(h_in)))


prs = Presentation(SRC)
print(f'Opened v32: {len(prs.slides)} slides')

# ── S51: head_angle_attribution_summary (no #1 annotation) ───────────────────
idx = find_idx(prs, 'Head Angle Is the Top-Attributed')
if idx is not None:
    replace_picture(prs.slides[idx],
                    os.path.join(mdir, 'head_angle_attribution_summary.png'),
                    0.25, 0.60, 9.50, 3.50)
    print(f'  S{idx+1}: replaced head_angle_attribution_summary.png')

# ── S60: position_ablation (Full bar now grey) ────────────────────────────────
idx = find_idx(prs, 'Position-Only MLP Recovers')
if idx is not None:
    replace_picture(prs.slides[idx],
                    os.path.join(mdir, 'position_ablation.png'),
                    0.25, 0.62, 9.50, 4.20)
    print(f'  S{idx+1}: replaced position_ablation.png')

# ── Verify S50 is untouched ───────────────────────────────────────────────────
idx50 = find_idx(prs, 'Head Angle Analysis')
if idx50 is not None:
    sl = prs.slides[idx50]
    texts = [sh.text_frame.text.strip() for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
    print(f'  S{idx50+1} (Head Angle title) retained: {texts}')

print(f'\nFinal slide count: {len(prs.slides)}')
for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved → {out}')
