#!/usr/bin/env python3
"""
build_v39.py  —  Insert head angle tuning variety slide between S53 and S54.

Base: user's Desktop ultimate_presentation_v38.pptx
v38 (92) → v39 (93 slides, +1):
  Insert after S53 "Head Angle MLP Training Evolution":
    "Head angle tuning: linear to non-monotonic"
    Figure: head_angle_tuning_variety.png  (9.50" × 4.20")
"""
import os, zipfile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.parts.presentation import PresentationPart


@property
def _safe_next_slide_partname(self):
    return self.package.next_partname('/ppt/slides/slide%d.xml')

PresentationPart._next_slide_partname = _safe_next_slide_partname

base = os.path.dirname(os.path.abspath(__file__))
root = os.path.join(base, '..')
mdir = os.path.join(root, 'outputs', 'mlps', 'ensembles_multiseed')

SRC  = '/mnt/c/Users/amits/Desktop/ultimate_presentation_v38.pptx'
OUTS = [
    os.path.join(root, 'outputs', 'ultimate_presentation_v39.pptx'),
    '/mnt/c/Users/amits/Desktop/ultimate_presentation_v39.pptx',
]

C_BG    = RGBColor(0xFF, 0xFF, 0xFF)
C_TITLE = RGBColor(0x1A, 0x23, 0x3A)
IMG     = os.path.join(mdir, 'head_angle_tuning_variety.png')
TITLE   = 'Head angle tuning: linear to non-monotonic'
TARGET_DPI = 200


def verify_figure(path, label):
    if not os.path.exists(path):
        print(f'  ✗ MISSING: {label}'); return False
    with Image.open(path) as im:
        w, h = im.size
    print(f'  ✓ {label}: {w}×{h}px = {w/TARGET_DPI:.2f}"×{h/TARGET_DPI:.2f}"')
    return True


def verify_zip(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        dupes = [n for n in set(names) if names.count(n) > 1]
        slides = [n for n in names if 'slides/slide' in n
                  and n.endswith('.xml') and 'rels' not in n]
    ok = len(dupes) == 0
    print(f'  {"✓" if ok else "✗ DUPES: "+str(dupes)} zip: {len(slides)} slides, {len(dupes)} dupes')
    return ok


def find_idx(prs, fragment):
    for i, sl in enumerate(prs.slides):
        for sh in sl.shapes:
            if sh.has_text_frame and fragment.lower() in sh.text_frame.text.lower():
                return i
    return None


def move_slide(prs, from_idx, to_idx):
    xml = prs.slides._sldIdLst
    els = list(xml); el = els.pop(from_idx); els.insert(to_idx, el)
    for c in list(xml): xml.remove(c)
    for c in els:        xml.append(c)


# ── Verify ────────────────────────────────────────────────────────────────────
print('PRE-BUILD VERIFICATION')
if not verify_figure(IMG, 'tuning variety'):
    raise SystemExit('Missing figure')
assert len(TITLE) <= 75, f'Title too long: {len(TITLE)} chars'
print(f'  ✓ title: {len(TITLE)} chars')
if not os.path.exists(SRC):
    raise SystemExit(f'Source not found: {SRC}')
print(f'  ✓ source: {os.path.basename(SRC)}\n')

# ── Build ─────────────────────────────────────────────────────────────────────
prs = Presentation(SRC)
print(f'Opened: {len(prs.slides)} slides')

# Find anchor: "Head Angle MLP Training Evolution" (S53)
anchor_idx = find_idx(prs, 'Head Angle MLP Training Evolution')
if anchor_idx is None:
    anchor_idx = find_idx(prs, 'Training Evolution')
print(f'  Anchor: S{anchor_idx+1} "{prs.slides[anchor_idx].shapes[0].text_frame.text[:50]}"')

# Build new slide
slide = prs.slides.add_slide(prs.slide_layouts[6])
fill  = slide.background.fill; fill.solid(); fill.fore_color.rgb = C_BG
tb = slide.shapes.add_textbox(Inches(0.20), Inches(0.06), Inches(9.60), Inches(0.52))
p  = tb.text_frame.paragraphs[0]; r = p.add_run()
r.text = TITLE
r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = C_TITLE
slide.shapes.add_picture(IMG,
    int(Inches(0.25)), int(Inches(0.62)),
    int(Inches(9.50)), int(Inches(4.20)))

new_idx    = len(prs.slides) - 1
insert_at  = anchor_idx + 1
move_slide(prs, new_idx, insert_at)
print(f'  Inserted at S{insert_at+1}')

# Context
print('\nContext:')
for i in range(anchor_idx, min(len(prs.slides), anchor_idx + 4)):
    sl = prs.slides[i]
    texts = [sh.text_frame.text.strip()[:65] for sh in sl.shapes if sh.has_text_frame]
    print(f'  S{i+1:02d}: {next((t for t in texts if t), "(empty)")}')

print(f'\nFinal: {len(prs.slides)} slides')

for out in OUTS:
    os.makedirs(os.path.dirname(out) or '.', exist_ok=True)
    prs.save(out)
    print(f'Saved → {out}')

print('\nPOST-BUILD VERIFICATION')
for out in OUTS:
    if os.path.exists(out):
        verify_zip(out)
